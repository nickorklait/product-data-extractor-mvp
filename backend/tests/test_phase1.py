import unittest

from openpyxl import load_workbook

from backend.app.excel import (
    create_ipd_consolidated_excel,
    load_ipd_template_rows,
    populate_ipd_rows,
    split_temperature_range,
    storage_temperature_parts,
)
from backend.app.models import (
    Allergens,
    DocumentMetadata,
    ExtractionResult,
    KeyValueCandidate,
    Nutrition,
    PhysicalChemicalData,
    RawLine,
    ShelfLife,
    Storage,
)


def make_result(product_name: str = "Phase 1 Test") -> ExtractionResult:
    return ExtractionResult(
        metadata=DocumentMetadata(source_file=f"{product_name}.docx", product_name=product_name),
        nutrition=Nutrition(energy_kj="1500", fat_g="12"),
        allergens=Allergens(gluten="No", milk="Yes", nuts="No"),
        shelf_life=ShelfLife(
            shelf_life_value="12",
            shelf_life_unit="måneder",
            shelf_life_text="Best before: 12 måneder from production date.",
        ),
        storage=Storage(
            storage_temperature="5-25 C",
            storage_text="Store dry. Temperature 5-25 °C.",
        ),
        physical_chemical=PhysicalChemicalData(density="1.2 g/ml"),
        key_value_candidates=[
            KeyValueCandidate(key="Vegan", value="Yes"),
            KeyValueCandidate(key="Halal", value="No"),
            KeyValueCandidate(key="Monounsaturates", value="3 g"),
            KeyValueCandidate(key="T", value="N E certificate noise"),
        ],
        raw_lines=[
            RawLine(line_number=1, section="allergens", text="Allergen value - gluten: No", source_reference="Page 2"),
            RawLine(line_number=2, section="allergens", text="Allergen value - nuts: No", source_reference="Page 2"),
        ],
    )


class Phase1SchemaTests(unittest.TestCase):
    def test_customer_template_contains_expected_attributes(self) -> None:
        rows = load_ipd_template_rows()

        self.assertEqual(len(rows), 52)
        self.assertEqual(rows[0]["Node"], "Allergens")
        self.assertEqual(rows[0]["Attribute"], "Cereals containing gluten and products thereof")
        self.assertEqual(rows[-1]["Attribute"], "Description")

    def test_phase1_mapping_uses_explicit_values_and_ignores_short_ocr_keys(self) -> None:
        result = populate_ipd_rows(make_result())
        mapped = {(row.node, row.attribute): row.data for row in result.ipd_rows}

        self.assertEqual(len(result.ipd_rows), 52)
        self.assertEqual(mapped[("Allergens", "Cereals containing gluten and products thereof")], "No")
        self.assertEqual(mapped[("Allergens", "Milk and products thereof (including lactose)")], "Yes")
        self.assertEqual(mapped[("Allergens", "Pecan")], "No")
        self.assertEqual(mapped[("Diet", "Vegan")], "Yes")
        self.assertEqual(mapped[("Religious", "Halal")], "No")
        self.assertIsNone(mapped[("Diet", "Vegetarian")])
        self.assertEqual(mapped[("Nutrient", "Energy kJ")], "1500")
        self.assertEqual(mapped[("Nutrient", "Monounsaturates")], "3")
        self.assertEqual(mapped[("General", "Storage temperature min/max")], "5 - 25")
        shelf_life = next(row for row in result.ipd_rows if row.attribute == "Shelf life")
        self.assertEqual(shelf_life.data, "12")
        self.assertEqual(shelf_life.uom, "months")
        self.assertEqual(shelf_life.full_text, "Best before: 12 måneder from production date.")
        gluten = next(row for row in result.ipd_rows if row.attribute.startswith("Cereals containing gluten"))
        pecan = next(row for row in result.ipd_rows if row.attribute == "Pecan")
        self.assertEqual(gluten.source_reference, "Page 2")
        self.assertEqual(pecan.source_reference, "Page 2")

    def test_comparison_workbook_shows_values_and_missing_status(self) -> None:
        first = populate_ipd_rows(make_result("Product A"))
        second = populate_ipd_rows(make_result("Product B"))
        second.ipd_rows = [row.model_copy(update={"data": None}) for row in second.ipd_rows]

        workbook = load_workbook(create_ipd_consolidated_excel([first, second]), data_only=True)

        self.assertEqual(workbook.sheetnames[:3], ["Comparison", "Completeness", "Missing_Attributes"])
        comparison = workbook["Comparison"]
        completeness = workbook["Completeness"]
        self.assertEqual(comparison.max_row, 3)
        headers = {cell.value: cell.column for cell in comparison[1]}
        self.assertEqual(comparison.cell(2, headers["General | Shelf life"]).value, 12)
        self.assertEqual(comparison.cell(2, headers["General | Shelf life | UoM"]).value, "months")
        self.assertEqual(
            comparison.cell(2, headers["General | Shelf life | Full text"]).value,
            "Best before: 12 måneder from production date.",
        )
        self.assertEqual(comparison.cell(2, headers["General | Storage temperature min/max"]).value, "5 - 25")
        self.assertEqual(comparison.cell(2, headers["General | Storage temperature min/max | UoM"]).value, "C")
        self.assertEqual(completeness.cell(2, 9).value, "Found")
        self.assertEqual(completeness.cell(3, 9).value, "Missing")
        self.assertEqual(comparison.cell(3, 9).value, "Not found")
        self.assertGreater(workbook["Missing_Attributes"].max_row, 1)
        workbook.close()

    def test_temperature_ranges_keep_only_real_negative_signs(self) -> None:
        self.assertEqual(split_temperature_range("5-25 C"), ("5", "25"))
        self.assertEqual(split_temperature_range("-18 to -12 C"), ("-18", "-12"))

    def test_temperature_range_can_be_recovered_from_storage_text(self) -> None:
        result = make_result()
        result.storage.storage_temperature = "25 C"
        result.storage.storage_text = "Lagringsbetingelser: tørt. Temperatur 10 � 25 °C."

        self.assertEqual(storage_temperature_parts(result), ("10", "25"))


if __name__ == "__main__":
    unittest.main()
