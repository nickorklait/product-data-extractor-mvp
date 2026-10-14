import unittest

from openpyxl import load_workbook

from backend.app.excel import create_ipd_consolidated_excel, populate_ipd_rows
from backend.app.finished_goods import FINISHED_GOOD_FIELDS, FINISHED_GOOD_PROFILE
from backend.app.parser import extract_structured_data


class SnacksDocumentTests(unittest.TestCase):
    def test_finished_good_schema_exports_one_product_per_row(self) -> None:
        text = """
Produkt | KiMs Test Chips
Nettovekt | 160 g
Materialnummer ferdigvare | 5958000001
Allergener | Melk
Claims | 40 % mindre fett
Flushing | Ja
Pakkegass | Nitrogen
Restoksygen | < 3,5 %
Materialnummer folie | 5928012509
EAN-kode kartong (D-pack) | 7039318018349
Antall poser per kartong | 25
Datomerking pose | PD + 180 dager
NÆRINGSINNHOLD Per 100 g
Energi (kJ/kcal) 1957/467
Fett 18 g
hvorav mettede fettsyrer 1,6 g
Karbohydrater 68 g
Protein 5,6 g
Salt 2,3 g
Jern | 4.2 mg | 1.3 mg | 10 %
"""

        result = extract_structured_data(text, "finished-good.docx")
        fields = {field.label: field for field in result.dynamic_fields}

        self.assertEqual(result.output_profile, FINISHED_GOOD_PROFILE)
        self.assertEqual(len(result.dynamic_fields), len(FINISHED_GOOD_FIELDS))
        self.assertEqual(fields["Produkt"].value, "KiMs Test Chips")
        self.assertEqual(fields["Pakkegass"].value, "Nitrogen")
        self.assertEqual(fields["Energi (kJ)"].value, "1957")
        self.assertEqual(fields["Energi (kJ)"].unit, "kJ")
        self.assertEqual(fields["Jern"].value, "4.2")
        self.assertEqual(fields["Jern"].unit, "mg")
        self.assertIsNone(fields["Magnesium"].value)

        workbook = load_workbook(create_ipd_consolidated_excel([result]), data_only=True)
        self.assertEqual(workbook.sheetnames[0], "Finished_Goods")
        sheet = workbook["Finished_Goods"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(2, headers["Produkt"]).value, "KiMs Test Chips")
        self.assertEqual(sheet.cell(2, headers["Energi (kJ)"]).value, 1957)
        self.assertEqual(sheet.cell(2, headers["Energi (kJ) UoM"]).value, "kJ")
        self.assertEqual(sheet.cell(2, headers["Jern"]).value, 4.2)
        self.assertEqual(sheet.cell(2, headers["Jern UoM"]).value, "mg")
        status_sheet = workbook["Finished_Goods_Status"]
        status_headers = {cell.value: cell.column for cell in status_sheet[1]}
        self.assertEqual(status_sheet.cell(2, status_headers["Jern"]).value, "Found")
        self.assertEqual(status_sheet.cell(2, status_headers["Magnesium"]).value, "Missing")
        self.assertEqual(workbook["Finished_Goods_Evidence"].max_row, len(FINISHED_GOOD_FIELDS) + 1)
        workbook.close()

    def test_pipe_nutrition_and_named_allergens_are_mapped(self) -> None:
        text = """
Produkt | Anyday Dry Roasted Nut Mix
Allergener | Produktet inneholder cashew, peanøtter, mandler og hasselnøtter. Kan inneholde spor av andre nøtter.
Datomerking pose | Produksjonsdato + 270 dager.
Næringsinnhold | Per 100 g | Per porsjon (30g)
Energi (kJ/kcal) | 2536/605 | 762/182
Fett | 49 g | 15 g
hvorav mettede fettsyrer | 7.1 g | 2.1 g
Karbohydrater | 18 g | 5.5 g
hvorav sukkerarter | 4.5 g | 1.3 g
Kostfiber | 7.5 g | 2.2 g
Protein | 20 g | 6.0 g
Salt | 0.8 g | 0.2 g
"""

        result = populate_ipd_rows(extract_structured_data(text, "anyday.docx"))
        mapped = {row.attribute: row.data for row in result.ipd_rows}

        self.assertEqual(result.nutrition.energy_kj, "2536")
        self.assertEqual(result.nutrition.energy_kcal, "605")
        self.assertEqual(result.nutrition.saturated_fat_g, "7.1")
        self.assertEqual(result.shelf_life.shelf_life_value, "270")
        self.assertEqual(result.allergens.peanuts, "Yes")
        self.assertEqual(result.allergens.nuts, "Yes")
        self.assertIn("andre nøtter", result.allergens.allergens_may_contain or "")
        self.assertEqual(mapped["Cashew"], "Yes")
        self.assertEqual(mapped["Hazelnut"], "Yes")
        self.assertEqual(mapped["Almonds"], "Yes")
        self.assertEqual(mapped["Other nuts"], "Traces")

    def test_finished_product_shelf_life_wins_and_process_temperature_is_ignored(self) -> None:
        text = """
Produkt | KiMs Snack Chips
Allergener | Melk
Holdbarhet pellets | 1 år
Veiledende Set.temp | 199 °C
Datomerking pose | PD + 180 dager
NÆRINGSINNHOLD Per 100 g 30g
Energi (kJ/kcal) 1957/467 581/139
Fett 18g 5,4g
hvorav mettede fettsyrer 1,6g 0,5g
Karbohydrat 68g 20g
hvorav sukkerarter 219 0,6g
Kostfiber 50g 15g
Protein 569 1,7g
Salt 239 0,7g
"""

        result = extract_structured_data(text, "snack-chips.docx")

        self.assertEqual(result.shelf_life.shelf_life_value, "180")
        self.assertEqual(result.shelf_life.shelf_life_unit, "dager")
        self.assertIsNone(result.storage.storage_temperature)
        self.assertEqual(result.allergens.milk, "Yes")
        self.assertEqual(result.nutrition.fat_g, "18")
        self.assertEqual(result.nutrition.carbohydrate_g, "68")
        self.assertIsNone(result.nutrition.sugars_g)
        self.assertIsNone(result.nutrition.fibre_g)
        self.assertIsNone(result.nutrition.protein_g)
        self.assertIsNone(result.nutrition.salt_g)

    def test_fragmented_nutrition_keeps_only_values_with_evidence(self) -> None:
        text = """
Produkt | IKEA Cashews
Allergener | Cashew, kan inneholde spor av andre nøtter.
£ Storage: Store atreomteme
Oppbevaring: Oppbevares ved romtemperatur.
word/media/image8.png:
PD+1 år Datomerkes under forhåndstrykket "best før".
Næringsinnhold pr. 100 g
Energi / Energi 2494 kJ
597 kcal
Fedt / Fett
heraf mættede fedtsyrer / hvorav mettede fettsyrer
Kulhydrat / Karbohydrater
heraf sukkerarter / hvorav sukkerarter
Protein / Protein 21g
Salt / Salt 0.30g
"""

        result = populate_ipd_rows(extract_structured_data(text, "ikea-cashews.docx"))
        mapped = {row.attribute: row.data for row in result.ipd_rows}

        self.assertEqual(result.nutrition.energy_kj, "2494")
        self.assertEqual(result.nutrition.energy_kcal, "597")
        self.assertEqual(result.nutrition.protein_g, "21")
        self.assertEqual(result.nutrition.salt_g, "0.30")
        self.assertIsNone(result.nutrition.fat_g)
        self.assertIsNone(result.nutrition.carbohydrate_g)
        self.assertEqual(mapped["Cashew"], "Yes")
        self.assertEqual(mapped["Other nuts"], "Traces")
        self.assertEqual(result.shelf_life.shelf_life_text, 'PD+1 år Datomerkes under forhåndstrykket "best før".')
        self.assertIn("Oppbevares ved romtemperatur", result.storage.storage_text or "")

    def test_no_allergens_sets_explicit_no_values(self) -> None:
        result = populate_ipd_rows(
            extract_structured_data("Produkt | Chips\nAllergener | Ingen allergener", "chips.docx")
        )
        allergen_rows = [row for row in result.ipd_rows if row.node == "Allergens"]

        self.assertTrue(allergen_rows)
        self.assertTrue(all(row.data == "No" for row in allergen_rows))


if __name__ == "__main__":
    unittest.main()
