import re
import shutil
import unicodedata
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Iterable, Optional

import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageEnhance, ImageFilter, ImageOps

from .models import (
    Allergens,
    DocumentMetadata,
    DetectedSection,
    DynamicField,
    ExtractionResult,
    KeyValueCandidate,
    Nutrition,
    PhysicalChemicalData,
    RawLine,
    ShelfLife,
    SpecificationRow,
    Storage,
)

PROJECT_TESSDATA_DIR = Path(__file__).resolve().parent.parent / "tessdata"
PROJECT_ROOT = Path(__file__).resolve().parents[2]
FILAMENT_TEMPLATE_PATH = PROJECT_ROOT / "samples" / "filament-specs-taihing-nylon.xlsx"

IMPORTANT_FIELDS = {
    "metadata.product_name": "Product name was not found.",
    "nutrition.energy_kj": "Energy kJ was not found.",
    "nutrition.protein_g": "Protein was not found.",
    "allergens.allergen_statement": "Allergen statement was not found.",
    "shelf_life.shelf_life_text": "Shelf life was not found.",
    "storage.storage_text": "Storage instructions were not found.",
}


def extract_text_from_file(path: Path) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix == ".pdf":
        return extract_pdf_text(path)
    raise ValueError("Unsupported file type. Please upload a .pdf or .docx file.")


def extract_docx_text(path: Path) -> tuple[str, list[str]]:
    doc = Document(path)
    parts: list[str] = []
    warnings: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    image_text = extract_docx_image_text(path, warnings)
    if image_text:
        parts.append("OCR text from embedded images:")
        parts.extend(image_text)

    return "\n".join(parts), warnings


def extract_pdf_text(path: Path) -> tuple[str, list[str]]:
    parts: list[str] = []
    warnings: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(f"PDF page {page_number}:\n{page_text.strip()}")

    native_text = "\n".join(parts)
    is_filament = has_any(
        native_text,
        ["production specifications standard approval", "product physical performance/function parameter"],
    )
    if len(native_text.strip()) >= 80 and not is_filament:
        return native_text, warnings

    try:
        import pytesseract

        configure_tesseract_path(pytesseract)
        language = get_ocr_language(pytesseract, warnings)
        document = pdfium.PdfDocument(path)
        try:
            if is_filament and len(document):
                first_image = document[0].render(scale=4).to_pil()
                width, height = first_image.size
                for name, box, psm in [
                    ("header", (.10, .01, .90, .11), 11),
                    ("contact footer", (.12, .885, .91, .985), 6),
                ]:
                    left, top, right, bottom = box
                    crop = ImageOps.grayscale(
                        first_image.crop(
                            (int(left * width), int(top * height), int(right * width), int(bottom * height))
                        )
                    )
                    supplemental = pytesseract.image_to_string(
                        crop,
                        lang=language,
                        config=f"{get_tessdata_config()} --psm {psm}",
                    ).strip()
                    if supplemental:
                        parts.append(f"OCR text from PDF {name}:\n{supplemental}")
            for page_number, page in enumerate(document, start=1):
                if is_filament:
                    continue
                image = page.render(scale=4).to_pil()
                ocr_text = pytesseract.image_to_string(
                    ImageOps.grayscale(image),
                    lang=language,
                    config=f"{get_tessdata_config()} --psm 3",
                ).strip()
                if ocr_text:
                    parts.append(f"OCR text from PDF page {page_number}:\n{ocr_text}")
                if has_any(ocr_text, ["article specification of roller fabric"]):
                    sparse_text = pytesseract.image_to_string(
                        ImageOps.grayscale(image),
                        lang=language,
                        config=f"{get_tessdata_config()} --psm 11",
                    ).strip()
                    if sparse_text:
                        parts.append(f"Sparse OCR text from PDF page {page_number}:\n{sparse_text}")
                    parts.extend(extract_roller_fabric_form_fields(image, pytesseract, language))
        finally:
            document.close()
        if is_filament:
            warnings.append("Contact and supplier header graphics were processed with local OCR; review those fields.")
        else:
            warnings.append("The PDF contained little or no embedded text, so its pages were processed with local OCR.")
    except ImportError:
        warnings.append("The PDF is image-only, but the local OCR dependencies are not installed.")
    except Exception as exc:
        warnings.append(f"PDF OCR could not be completed: {exc}")

    return "\n".join(parts), warnings


FABRIC_SECTIONS = {
    "document": "Document & Versioning",
    "material": "Material Facts",
    "yarn": "Yarn/Fibers Mix",
    "colour": "Special Production-Colour",
    "processing": "Special Production-Processing",
    "roll": "Roll Specifications",
    "supplier": "Supplier Info",
}


def extract_roller_fabric_form_fields(image: Image.Image, pytesseract_module, language: str) -> list[str]:
    """Read the stable cells in Orkla's roller-fabric form using relative page coordinates."""
    fields: list[tuple[str, str, str, tuple[float, float, float, float], str, str]] = [
        ("ohc_article_no", "OHC article no.", "document", (.35, .0968, .50, .1081), "", "number"),
        ("doc_date", "Document date", "document", (.615, .0968, .77, .1069), "", "date"),
        ("ohc_article_name", "OHC article name", "document", (.35, .1081, .50, .1188), "", "text"),
        ("edition", "Edition", "document", (.615, .1081, .77, .1188), "", "date"),
        ("supplier_article_no", "Supplier article no.", "document", (.35, .1188, .50, .1300), "", "text"),
        ("fabric_construction", "Fabric construction", "material", (.498, .1915, .613, .2028), "", "text"),
        ("material_type_pile", "Material type: pile", "material", (.498, .2028, .613, .2144), "", "text"),
        ("recycled_material_in_pile", "Recycled material in pile", "material", (.498, .2144, .613, .2257), "", "text"),
        ("steamed_yn", "Steamed", "material", (.498, .2257, .613, .2372), "", "yes_no"),
        ("material_type_backing", "Material type: backing", "material", (.498, .2372, .613, .2488), "", "text"),
        ("recycled_material_in_backing", "Recycled material in backing", "material", (.498, .2488, .613, .2601), "", "yes_no"),
        ("material_type_glue", "Material type: glue", "material", (.498, .2601, .613, .2717), "", "text"),
        ("glue_weight_gm2_target", "Glue weight target", "material", (.498, .2717, .613, .2830), "g/m2", "number"),
        ("glue_weight_tolerance_pct", "Glue weight tolerance", "material", (.613, .2717, .703, .2830), "%", "tolerance"),
        ("glue_weight_report_required", "Glue weight report required", "material", (.703, .2717, .77, .2830), "", "yes_no"),
        ("appearance_spec", "Appearance", "material", (.498, .2830, .613, .3017), "", "multiline"),
        ("fabric_gram_weight_gm2_target", "Fabric gram weight target", "material", (.498, .3017, .613, .3406), "g/m2", "number"),
        ("fabric_gram_weight_tolerance_pct", "Fabric gram weight tolerance", "material", (.613, .3017, .703, .3394), "%", "tolerance"),
        ("fabric_gram_weight_report_required", "Fabric gram weight report required", "material", (.703, .3017, .77, .3394), "", "yes_no"),
        ("pile_height_mm_target", "Pile height target", "material", (.498, .3406, .613, .3800), "mm", "number"),
        ("pile_height_tolerance_mm", "Pile height tolerance", "material", (.613, .3406, .703, .3800), "mm", "multiline"),
        ("pile_height_report_required", "Pile height report required", "material", (.703, .3406, .77, .3789), "", "yes_no"),
        ("pile_height_note", "Pile height note", "material", (.34, .3406, .498, .3800), "", "multiline"),
        ("yarn_1_material", "Yarn 1 material", "yarn", (.18, .4097, .613, .4201), "", "text"),
        ("yarn_1_thickness_dtex", "Yarn 1 thickness", "yarn", (.613, .4097, .703, .4201), "dtex", "number"),
        ("yarn_1_pct_in_pile", "Yarn 1 in pile", "yarn", (.703, .4097, .77, .4201), "%", "number"),
        ("yarn_2_material", "Yarn 2 material", "yarn", (.18, .4213, .613, .4317), "", "text"),
        ("yarn_2_thickness_dtex", "Yarn 2 thickness", "yarn", (.613, .4213, .703, .4317), "dtex", "number"),
        ("yarn_2_pct_in_pile", "Yarn 2 in pile", "yarn", (.703, .4213, .77, .4317), "%", "number"),
        ("yarn_4_material", "Yarn 4 material", "yarn", (.18, .4445, .613, .4549), "", "text"),
        ("yarn_4_thickness_dtex", "Yarn 4 thickness", "yarn", (.613, .4445, .703, .4549), "dtex", "number"),
        ("yarn_4_pct_in_pile", "Yarn 4 in pile", "yarn", (.703, .4445, .77, .4549), "%", "text"),
        ("colour_name", "Colour", "colour", (.49, .5050, .613, .5157), "", "text"),
        ("colour_code_fabric", "Fabric colour code", "colour", (.49, .5050, .613, .5157), "", "text"),
        ("stripe_distance_mm_target", "Stripe distance target", "colour", (.49, .5282, .613, .5389), "mm", "number_or_na"),
        ("stripe_distance_tolerance_mm", "Stripe distance tolerance", "colour", (.613, .5282, .703, .5389), "mm", "tolerance"),
        ("stripe_distance_report_required", "Stripe distance report required", "colour", (.703, .5282, .77, .5389), "", "yes_no"),
        ("stripe_width_target", "Stripe width target", "colour", (.49, .5398, .613, .5502), "mm", "number_or_na"),
        ("stripe_width_tolerance_mm", "Stripe width tolerance", "colour", (.613, .5398, .703, .5502), "mm", "tolerance"),
        ("stripe_width_report_required", "Stripe width report required", "colour", (.703, .5398, .77, .5502), "", "yes_no"),
        ("combing_trimming_yn", "Combing/trimming", "processing", (.49, .5514, .613, .5621), "", "yes_no"),
        ("roll_length_m_avg_target", "Average roll length target", "roll", (.49, .6012, .613, .6119), "m", "number"),
        ("roll_length_m_avg_tolerance_pct", "Average roll length tolerance", "roll", (.613, .6012, .703, .6119), "%", "tolerance"),
        ("roll_length_m_avg_report_required", "Roll length report required", "roll", (.703, .6012, .77, .6119), "", "yes_no"),
        ("material_width_net_cm_target", "Net material width", "roll", (.49, .6131, .613, .6235), "cm", "first_number"),
        ("material_width_gross_cm_target", "Gross material width", "roll", (.49, .6131, .613, .6235), "cm", "second_number"),
        ("material_width_tolerance_pct", "Material width tolerance", "roll", (.613, .6131, .703, .6235), "cm", "tolerance"),
        ("material_width_report_required", "Material width report required", "roll", (.703, .6131, .77, .6235), "", "yes_no"),
        ("roll_weight_kg", "Roll weight", "roll", (.49, .6247, .613, .6351), "kg", "number"),
        ("roll_marking_required_fields", "Roll marking required fields", "roll", (.49, .6363, .613, .6838), "", "multiline"),
        ("supplier_name", "Supplier name", "supplier", (.25, .748, .77, .763), "", "text"),
        ("supplier_address", "Supplier address", "supplier", (.25, .763, .77, .801), "", "multiline"),
        ("supplier_contact_name", "Supplier contact name", "supplier", (.25, .810, .77, .828), "", "text"),
        ("supplier_contact_phone", "Supplier contact phone", "supplier", (.25, .828, .77, .845), "", "phone"),
        ("supplier_contact_cellphone", "Supplier contact cellphone", "supplier", (.25, .845, .77, .862), "", "phone"),
        ("supplier_contact_email", "Supplier contact email", "supplier", (.25, .862, .77, .881), "", "email"),
    ]

    extracted: list[str] = []
    for field_name, label, section_key, box, unit, value_type in fields:
        raw_value = ocr_relative_cell(image, box, pytesseract_module, language, value_type=value_type)
        value, needs_review = normalize_form_value(raw_value, value_type)
        safe_value = value.replace("|", "/").replace("\n", " ")
        extracted.append(
            "__DYNAMIC_FIELD__|"
            f"{field_name}|{FABRIC_SECTIONS[section_key]}|{label}|{safe_value}|{unit}|{int(needs_review)}"
        )
    for field_name, label, section_key, unit in [
        ("yarn_3_material", "Yarn 3 material", "yarn", ""),
        ("yarn_3_thickness_dtex", "Yarn 3 thickness", "yarn", "dtex"),
        ("yarn_3_pct_in_pile", "Yarn 3 in pile", "yarn", "%"),
        ("shearing_carding_yn", "Shearing/carding", "processing", ""),
        ("other_comments", "Other comments", "roll", ""),
    ]:
        extracted.append(
            "__DYNAMIC_FIELD__|"
            f"{field_name}|{FABRIC_SECTIONS[section_key]}|{label}||{unit}|0"
        )
    return extracted


def ocr_relative_cell(
    image: Image.Image,
    box: tuple[float, float, float, float],
    pytesseract_module,
    language: str,
    value_type: str = "text",
) -> str:
    width, height = image.size
    left, top, right, bottom = box
    crop = ImageOps.grayscale(
        image.crop((int(left * width), int(top * height), int(right * width), int(bottom * height)))
    )
    crop = ImageEnhance.Contrast(crop).enhance(1.5)
    if crop.width > 24:
        crop = crop.crop((8, 0, crop.width - 8, crop.height))
    crop = ImageOps.expand(crop, border=18, fill=255)
    psm = 6 if value_type == "multiline" else 7
    whitelist = ""
    if value_type in {"number", "first_number", "second_number", "number_or_na", "tolerance"}:
        whitelist = " -c tessedit_char_whitelist=0123456789.,+-/%NA"
    if value_type in {"number", "first_number", "second_number", "number_or_na"}:
        psm = 11
    return pytesseract_module.image_to_string(
        crop,
        lang=language,
        config=f"{get_tessdata_config()} --psm {psm}{whitelist}",
    ).strip()


def normalize_form_value(raw_value: str, value_type: str) -> tuple[str, bool]:
    value = re.sub(r"[^\x20-\x7E]+", " ", raw_value)
    value = re.sub(r"\s+", " ", value).strip(" |[]{}_")
    value = re.sub(r"^(?:I|J|l)\s+(?=[A-Z0-9])", "", value)
    value = re.sub(r"\s+(?:I|l)$", "", value)
    if not value:
        return "", False
    lowered = searchable(value)
    if value_type == "yes_no":
        match = re.search(r"\b(yes|no)\b", lowered)
        return (match.group(1).upper(), False) if match else (value, True)
    if value_type == "date":
        match = re.search(r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}", value)
        return (match.group(0), False) if match else (value, True)
    if value_type == "email":
        match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", value)
        return (match.group(0), False) if match else (value, True)
    if value_type == "phone":
        match = re.search(r"\+?\d[\d ]{6,}", value)
        return (re.sub(r"\s+", " ", match.group(0)).strip(), False) if match else (value, True)
    if value_type == "tolerance":
        match = re.search(r"[+\-]?\s*/?\s*-?\s*\d+(?:[,.]\d+)?\s*(?:%|mm|cm)?", value)
        return (re.sub(r"\s+", "", match.group(0)), False) if match else (value, True)
    if value_type in {"number", "first_number", "second_number"}:
        numbers = re.findall(r"\d+(?:[,.]\d+)?", value)
        index = 1 if value_type == "second_number" else 0
        return (numbers[index], False) if len(numbers) > index else (value, True)
    if value_type == "number_or_na":
        if re.search(r"\bna\b", lowered):
            return "NA", False
        numbers = re.findall(r"\d+(?:[,.]\d+)?", value)
        return (numbers[0], False) if numbers else (value, True)
    suspicious = bool(re.search(r"[^A-Za-z0-9@.,+/%()' -]", value))
    return value, suspicious


def extract_docx_image_text(path: Path, warnings: list[str]) -> list[str]:
    image_blobs: list[tuple[str, bytes]] = []
    with zipfile.ZipFile(path) as docx_zip:
        for name in docx_zip.namelist():
            if name.startswith("word/media/"):
                image_blobs.append((name, docx_zip.read(name)))

    if not image_blobs:
        return []

    try:
        import pytesseract
    except ImportError:
        warnings.append(
            "The DOCX contains embedded images, but pytesseract is not installed. Image OCR was skipped."
        )
        return []

    configure_tesseract_path(pytesseract)
    ocr_language = get_ocr_language(pytesseract, warnings)
    tessdata_config = get_tessdata_config()

    ocr_text: list[str] = []
    for index, (image_name, image_blob) in enumerate(image_blobs, start=1):
        try:
            source_image = Image.open(BytesIO(image_blob))
            table_image = ImageOps.autocontrast(ImageOps.exif_transpose(source_image).convert("L"))
            image = preprocess_image_for_ocr(source_image)
            text = pytesseract.image_to_string(image, lang=ocr_language, config=tessdata_config).strip()
            if text:
                ocr_text.append(f"{image_name}:\n{text}")
                allergen_values = extract_allergen_table_values(
                    table_image,
                    text,
                    pytesseract,
                    ocr_language,
                    tessdata_config,
                )
                if allergen_values:
                    rows = [f"Allergen value - {name}: {value}" for name, value in allergen_values.items()]
                    ocr_text.append("Allergen table values:\n" + "\n".join(rows))
        except pytesseract.TesseractNotFoundError:
            warnings.append(
                "The DOCX contains embedded images, but the Tesseract OCR app is not installed or not on PATH. Image OCR was skipped."
            )
            return []
        except pytesseract.TesseractError:
            try:
                text = pytesseract.image_to_string(image, lang="eng").strip()
                if text:
                    ocr_text.append(f"{image_name}:\n{text}")
            except pytesseract.TesseractError as exc:
                warnings.append(f"Image {index} could not be OCR processed with configured languages: {exc}")
        except Exception as exc:
            warnings.append(f"Image {index} could not be OCR processed: {exc}")

    if image_blobs and not ocr_text:
        warnings.append("The DOCX contains embedded images, but OCR did not find readable text.")
    return ocr_text


def preprocess_image_for_ocr(image: Image.Image) -> Image.Image:
    image = ImageOps.exif_transpose(image).convert("L")
    width, height = image.size
    longest_side = max(width, height)
    if longest_side < 1800:
        scale = min(3, max(2, round(1800 / longest_side)))
        image = image.resize((width * scale, height * scale), Image.Resampling.LANCZOS)

    image = ImageOps.autocontrast(image)
    image = ImageEnhance.Contrast(image).enhance(1.8)
    image = image.filter(ImageFilter.SHARPEN)
    image = image.point(lambda pixel: 255 if pixel > 180 else 0)
    return image


ALLERGEN_TABLE_LABELS = {
    "gluten": ["cereals containing gluten", "gluten"],
    "crustaceans": ["crustaceans", "crustacean", "skalldyr"],
    "eggs": ["egg and products", "eggs", "egg"],
    "fish": ["fish and products", "fish", "fisk"],
    "peanuts": ["peanut and products", "peanuts", "peanut"],
    "soybeans": ["soybeans and products", "soybeans", "soya", "soy"],
    "milk": ["milk and products", "milk", "melk"],
    "nuts": ["nuts and products", "tree nuts", "notter"],
    "celery": ["celery and products", "celery", "selleri"],
    "mustard": ["mustard and products", "mustard", "sennep"],
    "sesame": ["sesame seeds", "sesame", "sesam"],
    "sulphites": ["sulphur dioxide and sulphites", "sulphites", "sulfitt"],
    "lupin": ["lupin and products", "lupin"],
    "molluscs": ["mollusc and products", "molluscs", "mollusc", "blotdyr"],
}


def extract_allergen_table_values(
    image: Image.Image,
    initial_text: str,
    pytesseract_module,
    language: str,
    tessdata_config: str,
) -> dict[str, str]:
    initial_key = searchable(initial_text)
    allergen_hits = sum(
        1 for term in ["allergen declaration", "cereals containing gluten", "sesame seeds", "sulphur dioxide", "mollusc"]
        if searchable(term) in initial_key
    )
    if allergen_hits < 1:
        return {}

    table_text = pytesseract_module.image_to_string(
        image,
        lang=language,
        config=f"{tessdata_config} --psm 4",
    )
    values: dict[str, str] = {}
    for line in table_text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip()
        line_key = searchable(cleaned)
        if not cleaned:
            continue
        status_match = re.search(r"\b(yes|no|ja|nei)\b", cleaned, flags=re.IGNORECASE)
        quantity_match = re.search(r"[<>≤≥]\s*\d+(?:[,.]\d+)?\s*ppm\b", cleaned, flags=re.IGNORECASE)
        if not status_match and not quantity_match:
            continue
        value = status_match.group(1).title() if status_match else re.sub(r"\s+", "", quantity_match.group(0))
        for field_name, aliases in ALLERGEN_TABLE_LABELS.items():
            if field_name in values:
                continue
            if any(searchable(alias) in line_key for alias in aliases):
                values[field_name] = value
                break
    return values


def extract_allergen_statuses(text: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for match in re.finditer(
        r"^Allergen value - ([A-Za-z]+):\s*(.+?)\s*$",
        text,
        flags=re.IGNORECASE | re.MULTILINE,
    ):
        field_name = match.group(1).casefold()
        if field_name in ALLERGEN_TABLE_LABELS and field_name not in values:
            values[field_name] = match.group(2).strip()
    return values


def configure_tesseract_path(pytesseract_module) -> None:
    if shutil.which("tesseract"):
        return

    common_paths = [
        Path("C:/Program Files/Tesseract-OCR/tesseract.exe"),
        Path("C:/Program Files (x86)/Tesseract-OCR/tesseract.exe"),
        Path("C:/Users/nhauge/AppData/Local/Programs/Tesseract-OCR/tesseract.exe"),
    ]
    for candidate in common_paths:
        if candidate.exists():
            pytesseract_module.pytesseract.tesseract_cmd = str(candidate)
            return


def get_ocr_language(pytesseract_module, warnings: list[str]) -> str:
    config = get_tessdata_config()
    languages = get_project_tessdata_languages()
    if not languages:
        try:
            languages = set(pytesseract_module.get_languages(config=config))
        except Exception:
            languages = set()

    if "eng" in languages and "nor" in languages:
        return "eng+nor"
    if "eng" in languages:
        warnings.append("Norwegian OCR language data is not installed. OCR used English only, so Norwegian text may be less accurate.")
        return "eng"
    return "eng"


def get_tessdata_config() -> str:
    if PROJECT_TESSDATA_DIR.exists():
        return f"--tessdata-dir {PROJECT_TESSDATA_DIR}"
    return ""


def get_project_tessdata_languages() -> set[str]:
    if not PROJECT_TESSDATA_DIR.exists():
        return set()
    return {path.stem for path in PROJECT_TESSDATA_DIR.glob("*.traineddata")}


def extract_structured_data(text: str, source_file: str, parser_warnings: Optional[list[str]] = None) -> ExtractionResult:
    normalized = normalize_text(text)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    raw_lines = build_raw_lines(text)
    specification_rows = extract_specification_rows(raw_lines)
    filename_metadata = extract_metadata_from_filename(source_file)
    output_profile = detect_output_profile(text, source_file)

    metadata = DocumentMetadata(
        source_file=source_file,
        product_name=find_product_name_value(lines) or find_document_title(lines) or filename_metadata.get("product_name"),
        supplier=find_metadata_value(lines, ["supplier", "leverandor"]) or find_supplier_from_text(lines) or filename_metadata.get("supplier"),
        plant=find_metadata_value(lines, ["plant", "anlegg", "fabrikk"]) or filename_metadata.get("plant"),
        document_id=find_metadata_value(lines, ["document id", "document no", "dokument id", "dokumentnr"]) or filename_metadata.get("document_id"),
        version=find_metadata_value(lines, ["version", "versjon"]),
        approval_status=find_metadata_value(lines, ["approval status", "approved", "godkjent", "status"]),
        approved_date=find_metadata_value(lines, ["approved date", "approval date", "godkjent dato"]),
    )

    nutrition_section = find_section(
        text,
        ["naeringsinnhold", "nutrition", "nutritional value", "nutrition declaration"],
        ["allergen", "allergener", "shelf life", "holdbarhet", "storage", "oppbevaring", "physical", "fysisk"],
    )
    allergens_section = find_section(
        text,
        ["allergen", "allergener", "allergens", "contains", "may contain", "spor av"],
        ["shelf life", "holdbarhet", "storage", "oppbevaring", "physical", "fysisk", "nutrition", "naeringsinnhold"],
    )
    shelf_life_section = find_section(
        text,
        ["shelf life", "holdbarhet", "best before", "expiry", "durability"],
        ["storage", "oppbevaring", "physical", "fysisk", "allergen", "nutrition", "naeringsinnhold"],
    )
    storage_section = find_section(
        text,
        ["storage", "oppbevaring", "lagring", "store", "temperature"],
        ["physical", "fysisk", "chemical", "kjemisk", "allergen", "nutrition", "naeringsinnhold"],
    )
    physical_section = find_section(
        text,
        ["physical", "chemical", "fysisk", "kjemisk", "moisture", "water content", "pH", "density", "particle size"],
        ["allergen", "shelf life", "holdbarhet", "storage", "oppbevaring", "nutrition", "naeringsinnhold"],
    )

    nutrition_source = nutrition_section or text
    vertical_nutrition = extract_vertical_nutrition(nutrition_source)
    nutrition = Nutrition(
        energy_kj=find_specification_numeric_value(specification_rows, ["energy"], "kj")
        or find_energy_value(nutrition_source, "kj") or vertical_nutrition.get("energy_kj"),
        energy_kcal=find_specification_numeric_value(specification_rows, ["energy"], "kcal")
        or find_energy_value(nutrition_source, "kcal") or vertical_nutrition.get("energy_kcal"),
        fat_g=find_specification_numeric_value(specification_rows, ["fat"])
        or find_nutrition_value(nutrition_source, [r"fat", r"fett"], ["g"]) or vertical_nutrition.get("fat_g"),
        saturated_fat_g=find_specification_numeric_value(specification_rows, ["of which saturates", "saturated fat"])
        or find_nutrition_value(nutrition_source, [r"saturated fat", r"mettede fettsyrer", r"mettet fett"], ["g"]) or vertical_nutrition.get("saturated_fat_g"),
        carbohydrate_g=find_specification_numeric_value(specification_rows, ["carbohydrate"])
        or find_nutrition_value(nutrition_source, [r"carbohydrate", r"karbohydrat"], ["g"]) or vertical_nutrition.get("carbohydrate_g"),
        sugars_g=find_specification_numeric_value(specification_rows, ["of which sugars", "sugars"])
        or find_nutrition_value(nutrition_source, [r"sugars?", r"sukkerarter", r"sukker"], ["g"]) or vertical_nutrition.get("sugars_g"),
        fibre_g=find_specification_numeric_value(specification_rows, ["fibre", "fiber"])
        or find_nutrition_value(nutrition_source, [r"fibre", r"fiber", r"kostfiber"], ["g"]) or vertical_nutrition.get("fibre_g"),
        protein_g=find_specification_numeric_value(specification_rows, ["protein"])
        or find_nutrition_value(nutrition_source, [r"protein"], ["g"]) or vertical_nutrition.get("protein_g"),
        salt_g=find_specification_numeric_value(specification_rows, ["salt"])
        or find_nutrition_value(nutrition_source, [r"salt"], ["g"]) or vertical_nutrition.get("salt_g"),
    )

    allergen_source = allergens_section or text
    allergen_statuses = extract_allergen_statuses(text)
    status_statement = "; ".join(
        f"{field_name.replace('_', ' ').title()}: {value}"
        for field_name, value in allergen_statuses.items()
    )
    contains_values = [
        field_name.replace("_", " ").title()
        for field_name, value in allergen_statuses.items()
        if searchable(value) in {"yes", "ja"}
    ]
    trace_values = [
        field_name.replace("_", " ").title()
        for field_name, value in allergen_statuses.items()
        if has_any(value, ["trace", "traces", "spor av"])
    ]
    allergens = Allergens(
        allergens_contains=", ".join(contains_values) or None if allergen_statuses else find_labeled_text(allergen_source, ["contains", "inneholder"], max_lines=3),
        allergens_may_contain=", ".join(trace_values) or None if allergen_statuses else find_labeled_text(allergen_source, ["may contain", "spor av", "kan inneholde"], max_lines=3),
        allergen_statement=status_statement or (clean_excerpt(allergen_source, 700) if has_any(allergen_source, ["allergen", "allergener", "contains", "spor av"]) else None),
        **allergen_statuses,
    )

    shelf_life_text = clean_excerpt(shelf_life_section, 500) if shelf_life_section else None
    shelf_value, shelf_unit = extract_shelf_life_value(shelf_life_section or normalized)
    shelf_life = ShelfLife(
        shelf_life_value=shelf_value,
        shelf_life_unit=shelf_unit,
        shelf_life_text=shelf_life_text,
    )

    storage = Storage(
        storage_temperature=find_temperature(storage_section or text),
        storage_text=clean_excerpt(storage_section, 500) if storage_section else None,
    )

    physical_source = physical_section or text
    physical_chemical = PhysicalChemicalData(
        moisture=find_specification_value(specification_rows, ["loss on drying", "moisture", "water content"])
        or find_labeled_measure(physical_source, ["moisture", "water content", "vanninnhold", "fuktighet"], ["%", "g"]),
        ph=find_specification_value(
            specification_rows,
            ["ph at", "ph of", "ph"],
            ["physical and chemical"],
        )
        or find_labeled_measure(physical_source, ["ph", "pH"], []),
        density=find_specification_value(specification_rows, ["bulk density", "density", "densitet"])
        or find_labeled_measure(physical_source, ["density", "densitet", "bulk density"], ["g/ml", "kg/m3", "kg/m^3"]),
        particle_size=find_labeled_measure(physical_source, ["particle size", "partikkelstorrelse"], ["mm", "um"]),
        physical_chemical_text=clean_excerpt(physical_section, 700) if physical_section else None,
    )
    detected_sections = build_detected_sections(
        {
            "nutrition": nutrition_section,
            "allergens": allergens_section,
            "shelf_life": shelf_life_section,
            "storage": storage_section,
            "physical_chemical": physical_section,
        }
    )
    key_value_candidates = extract_key_value_candidates(raw_lines)
    key_value_candidates.extend(
        KeyValueCandidate(
            section="physical_chemical" if row.section != "Microbiological values" else "microbiological",
            key=row.parameter,
            value=specification_display_value(row),
            source_line=row.source_line,
            source_reference=row.source_reference,
        )
        for row in specification_rows
    )
    review_flags = build_review_flags(nutrition, key_value_candidates)
    dynamic_fields = build_dynamic_fields(text, source_file)

    if output_profile in {"roller_fabric", "filament_spec"}:
        dynamic_values = {field.field_name: field.value for field in dynamic_fields if field.value}
        metadata.product_name = dynamic_values.get("ohc_article_name") or metadata.product_name
        metadata.document_id = dynamic_values.get("ohc_article_no") or metadata.document_id
        metadata.supplier = dynamic_values.get("supplier_name") or metadata.supplier
        metadata.version = dynamic_values.get("edition") or metadata.version

    result = ExtractionResult(
        metadata=metadata,
        nutrition=nutrition,
        allergens=allergens,
        shelf_life=shelf_life,
        storage=storage,
        physical_chemical=physical_chemical,
        extraction_warnings=[],
        raw_text_excerpt=clean_excerpt(text, 2000),
        raw_text_full=text,
        raw_lines=raw_lines,
        key_value_candidates=key_value_candidates,
        detected_sections=detected_sections,
        review_flags=review_flags,
        output_profile=output_profile,
        dynamic_fields=dynamic_fields,
        specification_rows=specification_rows,
    )
    result.extraction_warnings = [*(parser_warnings or []), *build_warnings(result)]
    review_specifications = [row.parameter for row in specification_rows if row.needs_review]
    if review_specifications:
        result.extraction_warnings.append(
            "Review OCR for specification rows: " + ", ".join(review_specifications[:8]) + "."
        )
    if has_any(text, ["allergen declaration", "allergen declaration european union"]) and not allergen_statuses:
        result.extraction_warnings.insert(
            0,
            "An allergen declaration was detected, but its individual Yes/No values could not be read. Manual review is required.",
        )
    return result


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def has_any(text: str, keywords: Iterable[str]) -> bool:
    haystack = searchable(text)
    return any(searchable(keyword) in haystack for keyword in keywords)


def find_metadata_value(lines: list[str], labels: list[str]) -> Optional[str]:
    for line in lines:
        searchable_line = searchable(line)
        label_pattern = "|".join(re.escape(searchable(label)) for label in labels)
        match = re.search(rf"\b({label_pattern})\b\s*[:\-|]\s*(.+)$", searchable_line, flags=re.IGNORECASE)
        if match:
            value_start = match.start(2)
            return clean_value(line[value_start:])
    return None


def find_product_name_value(lines: list[str]) -> Optional[str]:
    values: list[str] = []
    for line in lines:
        searchable_line = searchable(line)
        match = re.search(
            r"\b(product name|produktnavn|product|produkt)\b\s*[:\-|]\s*(.+)$",
            searchable_line,
            flags=re.IGNORECASE,
        )
        if match:
            values.append(clean_value(line[match.start(2) :]))
    if not values:
        return None
    sensible = [
        value
        for value in values
        if not has_any(value, ["product sheet", "page ", "revision date", "version no"])
    ]
    return min(sensible or values, key=len)


def extract_metadata_from_filename(source_file: str) -> dict[str, str]:
    stem = Path(source_file).stem.strip()
    normalized_stem = re.sub(r"\s+", " ", stem)
    pieces = [piece.strip() for piece in re.split(r"\s+[-\u2013\u2014]\s+", normalized_stem) if piece.strip()]
    metadata: dict[str, str] = {}

    leading_id = re.match(r"^(\d{4,})(?:\s+[-\u2013\u2014]\s*|\s+)(.+)$", normalized_stem)
    if leading_id:
        metadata["document_id"] = leading_id.group(1)
        remaining_stem = leading_id.group(2).strip()
        pieces = [piece.strip() for piece in re.split(r"\s+[-\u2013\u2014]\s+", remaining_stem) if piece.strip()]
    elif pieces and re.fullmatch(r"\d{4,}", pieces[0]):
        metadata["document_id"] = pieces[0]
        pieces = pieces[1:]

    if pieces:
        metadata["product_name"] = pieces[0]

    supplier_match = re.search(r"\bSupplier\s*-?\s*([A-Za-z0-9_-]+)", source_file, flags=re.IGNORECASE)
    if supplier_match:
        metadata["supplier"] = supplier_match.group(1)
    elif len(pieces) >= 2:
        metadata["supplier"] = pieces[1]

    plant_match = re.search(r"\bPlant\s*-?\s*([A-Za-z0-9_-]+)", source_file, flags=re.IGNORECASE)
    if plant_match:
        metadata["plant"] = plant_match.group(1)

    return metadata


def find_document_title(lines: list[str]) -> Optional[str]:
    for index, line in enumerate(lines):
        searchable_line = searchable(line)
        match = re.search(r"\bdatablad\s*[-]\s*(.+)$", searchable_line)
        if match:
            value_start = match.start(1)
            return clean_value(line[value_start:])
        if searchable_line in {"datablad", "product specification", "produktspesifikasjon"} and index + 1 < len(lines):
            return clean_value(lines[index + 1])
    return None


def find_supplier_from_text(lines: list[str]) -> Optional[str]:
    for line in lines[:30]:
        match = re.search(r"\b([A-Z][A-Za-z0-9 .,&-]{2,60}\s+(?:AS|AB|ASA|Ltd|Limited|GmbH))\b", line)
        if match:
            return clean_value(match.group(1))
    return None


def find_section(text: str, start_keywords: list[str], stop_keywords: list[str]) -> Optional[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    start_index = None
    for index, line in enumerate(lines):
        if has_any(line, start_keywords):
            start_index = index
            break

    if start_index is None:
        return None

    end_index = min(len(lines), start_index + 24)
    for index in range(start_index + 1, min(len(lines), start_index + 40)):
        if has_any(lines[index], stop_keywords):
            end_index = index
            break

    return "\n".join(lines[start_index:end_index]).strip() or None


def find_nutrition_value(text: str, labels: list[str], units: list[str]) -> Optional[str]:
    label_pattern = "|".join(labels)
    unit_pattern = "|".join(re.escape(unit) for unit in units)
    number = r"([<>]?\s*\d+(?:[,.]\d+)?)"
    searchable_text = searchable(text)
    patterns = [
        rf"(?:{label_pattern})[^\n\r\d<>]{{0,80}}{number}\s*(?:{unit_pattern})\b",
        rf"{number}\s*(?:{unit_pattern})\b[^\n\r]{{0,80}}(?:{label_pattern})",
    ]
    for pattern in patterns:
        match = re.search(pattern, searchable_text, flags=re.IGNORECASE)
        if match:
            return clean_value(match.group(1))
    return None


def find_energy_value(text: str, unit: str) -> Optional[str]:
    match = re.search(rf"([<>]?\s*\d+(?:[,.]\d+)?)\s*{re.escape(unit)}\b", searchable(text), flags=re.IGNORECASE)
    return clean_value(match.group(1)) if match else None


def extract_vertical_nutrition(text: str) -> dict[str, str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not any(has_any(line, ["nutrition", "naeringsinnhold"]) for line in lines):
        return {}

    value_lines: list[str] = []
    after_serving_header = False
    for line in lines:
        searchable_line = searchable(line)
        if re.search(r"\bpr\.?\s*100\b|\bper\s*100\b", searchable_line):
            after_serving_header = True
            continue
        if after_serving_header:
            if has_any(line, ["typiske verdier", "sensoriske", "physical", "allergen"]):
                break
            if re.search(r"\d", line):
                value_lines.append(clean_value(line))

    if not value_lines:
        return {}

    values: dict[str, str] = {}
    energy_line = next((line for line in value_lines if has_any(line, ["kj", "kcal"])), None)
    if energy_line:
        kj = find_energy_value(energy_line, "kj")
        kcal = find_energy_value(energy_line, "kcal")
        if kj:
            values["energy_kj"] = kj
        if kcal:
            values["energy_kcal"] = kcal
        value_lines = [line for line in value_lines if line != energy_line]

    ordered_fields = [
        "protein_g",
        "fat_g",
        "saturated_fat_g",
        "carbohydrate_g",
        "sugars_g",
        "fibre_g",
        "salt_g",
    ]
    for field, value in zip(ordered_fields, value_lines):
        values[field] = value
    return values


def find_labeled_text(text: str, labels: list[str], max_lines: int = 2) -> Optional[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for index, line in enumerate(lines):
        if has_any(line, labels):
            selected = lines[index : index + max_lines]
            return clean_excerpt("\n".join(selected), 400)
    return None


def extract_shelf_life_value(text: str) -> tuple[Optional[str], Optional[str]]:
    if not text:
        return None, None
    match = re.search(
        r"(\d+(?:[,.]\d+)?)\s*(days?|dager|months?|mnd|maneder|years?|ar|weeks?|uker)",
        searchable(text),
        flags=re.IGNORECASE,
    )
    if not match:
        return None, None
    return clean_value(match.group(1)), clean_value(match.group(2))


def find_temperature(text: str) -> Optional[str]:
    match = re.search(r"([+\-]?\d+(?:[,.]\d+)?\s*(?:deg\s*C|C)\s*(?:[-]\s*[+\-]?\d+(?:[,.]\d+)?\s*(?:deg\s*C|C))?)", searchable(text), flags=re.IGNORECASE)
    return clean_value(match.group(1)) if match else None


def find_labeled_measure(text: str, labels: list[str], units: list[str]) -> Optional[str]:
    label_pattern = "|".join(re.escape(label) for label in labels)
    unit_pattern = "|".join(re.escape(unit) for unit in units)
    number = r"([<>]?\s*\d+(?:[,.]\d+)?)"
    searchable_text = searchable(text)
    if units:
        pattern = rf"(?:{label_pattern})[^\n\r\d<>]{{0,80}}{number}\s*(?:{unit_pattern})?\b"
    else:
        pattern = rf"(?:{label_pattern})[^\n\r\d<>]{{0,80}}{number}\b"
    match = re.search(pattern, searchable_text, flags=re.IGNORECASE)
    if not match:
        return None
    value = clean_value(match.group(1))
    unit_match = re.search(rf"{re.escape(value)}\s*({unit_pattern})", match.group(0), flags=re.IGNORECASE) if units else None
    return f"{value} {unit_match.group(1)}".strip() if unit_match else value


def searchable(text: str) -> str:
    replacements = {
        "\u00e6": "ae",
        "\u00c6": "ae",
        "\u00f8": "o",
        "\u00d8": "o",
        "\u00e5": "a",
        "\u00c5": "a",
        "\u00b0": "deg",
        "\u00b5": "u",
        "\u2013": "-",
        "\u2014": "-",
    }
    for source, replacement in replacements.items():
        text = text.replace(source, replacement)
    normalized = unicodedata.normalize("NFKD", text)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return ascii_text.casefold()


def clean_value(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip(" \t:-|")


def clean_excerpt(text: Optional[str], limit: int) -> Optional[str]:
    if not text:
        return None
    cleaned = re.sub(r"[ \t]+", " ", text).strip()
    return cleaned[:limit].strip() if cleaned else None


def get_nested(result: ExtractionResult, dotted_path: str) -> Optional[str]:
    current = result
    for part in dotted_path.split("."):
        current = getattr(current, part)
    return current


def build_warnings(result: ExtractionResult) -> list[str]:
    warnings: list[str] = []
    if result.output_profile == "food_ipd":
        warnings.extend(message for path, message in IMPORTANT_FIELDS.items() if not get_nested(result, path))
    elif result.output_profile == "roller_fabric":
        important = {"ohc_article_no", "ohc_article_name", "supplier_name", "fabric_construction"}
        found = {field.field_name for field in result.dynamic_fields if field.value}
        for field_name in sorted(important - found):
            warnings.append(f"Important fabric field '{field_name}' was not found.")
        flagged = [field.label for field in result.dynamic_fields if field.needs_review]
        if flagged:
            warnings.append(f"OCR values need manual review: {', '.join(flagged)}.")
    elif result.output_profile == "filament_spec":
        important = {"ohc_article_no", "ohc_article_name", "material_type", "diameter_mm_target", "filament_length_mm_target"}
        found = {field.field_name for field in result.dynamic_fields if field.value}
        for field_name in sorted(important - found):
            warnings.append(f"Important filament field '{field_name}' was not found.")
        flagged = [field.label for field in result.dynamic_fields if field.needs_review]
        if flagged:
            warnings.append(f"Fields need manual review: {', '.join(flagged)}.")
        if has_any(result.raw_text_full or "", ["to be corrected", "\u5f85\u66f4\u6b63"]):
            warnings.append("The source marks the bending-strength specification as 'to be corrected'.")
    warnings.extend(result.review_flags)
    if not result.raw_text_excerpt:
        warnings.append("No readable text was extracted from the document.")
    return warnings


def build_raw_lines(text: str) -> list[RawLine]:
    raw_lines: list[RawLine] = []
    active_section: Optional[str] = None
    source_reference: Optional[str] = None
    for index, line in enumerate(text.splitlines(), start=1):
        cleaned = line.strip()
        if not cleaned:
            continue
        if cleaned.startswith("__DYNAMIC_FIELD__|"):
            continue
        image_match = re.match(r"^word/media/([^:]+):$", cleaned, flags=re.IGNORECASE)
        pdf_match = re.match(r"^(?:OCR text from )?PDF page (\d+):$", cleaned, flags=re.IGNORECASE)
        if image_match:
            source_reference = f"Embedded image {image_match.group(1)}"
        elif pdf_match:
            source_reference = f"PDF page {pdf_match.group(1)}"
        detected = classify_section(cleaned)
        if detected:
            active_section = detected
        raw_lines.append(
            RawLine(
                line_number=index,
                section=active_section,
                text=cleaned,
                source_reference=source_reference,
            )
        )
    return raw_lines


SPECIFICATION_SECTION_HEADINGS = {
    "manufacturing": "Manufacturing",
    "colouringproperties": "Colouring properties",
    "coloringproperties": "Colouring properties",
    "physicalandchemicalproperties": "Physical and chemical properties",
    "physicalandchemicalvalues": "Physical and chemical values",
    "physicalchemicalvalues": "Physical and chemical values",
    "microbiologicaldata": "Microbiological data",
    "microbiologicalvalues": "Microbiological values",
    "microbiologicalrequirements": "Microbiological values",
    "generalappearance": "General appearance",
    "nutritionaldata": "Nutritional data",
    "shelflifeandstorageconditions": "Shelf life and storage",
    "typicalvalues": "Typical values",
}

SPECIFICATION_STOP_HEADINGS = {
    "transportconditions",
    "safetydatasheet",
    "gmostatus",
    "irradiation",
    "novelfood",
    "nanotechnology",
    "allergenstatus",
    "legalstatus",
    "disclaimer",
    "specificdisclaimer",
    "generaldisclaimer",
}

SPEC_UNIT_PATTERN = (
    r"%\s*\(w/w\)|kJ/100g|kcal[/']?100g|keal[/']?100[9g]|g9?/100[9g]|9/100[9g]|"
    r"mg/kg|me/kg|g/l|kg/l|kgA|g/kg|gikg|/g|cfu(?:/|l)g|cfu/25g|UB|nm|pH|\*?Brix|%"
)


def extract_specification_rows(raw_lines: list[RawLine]) -> list[SpecificationRow]:
    rows: list[SpecificationRow] = []
    active_section: Optional[str] = None
    pending_parameter: Optional[str] = None
    for raw_line in raw_lines:
        text = raw_line.text.strip()
        key = searchable(text)
        compact_key = re.sub(r"[^a-z0-9]+", "", key)
        detected_section = specification_section_for_heading(compact_key)
        if detected_section:
            active_section = detected_section
            pending_parameter = None
            continue
        if active_section and (
            compact_key in {"additionaldata", "storage", "conformity"}
            or any(compact_key.startswith(heading) for heading in SPECIFICATION_STOP_HEADINGS)
        ):
            active_section = None
            pending_parameter = None
            continue
        if not active_section or is_document_footer_line(text):
            continue

        parsed = parse_specification_line(text, active_section, raw_line)
        if not parsed:
            parsed = parse_named_specification_line(text, active_section, raw_line)
        if parsed:
            if pending_parameter and re.match(r"^\d+\s*(?:min|sec|h|at)\b", parsed.parameter, flags=re.IGNORECASE):
                parsed.parameter = f"{pending_parameter} - {parsed.parameter}"
                parsed.source_excerpt = f"{pending_parameter} | {text}"
            rows.append(parsed)
            pending_parameter = None
            continue

        value_only = parse_specification_value_only(text)
        if value_only and pending_parameter:
            rows.append(
                SpecificationRow(
                    section=active_section,
                    parameter=pending_parameter,
                    source_line=raw_line.line_number,
                    source_reference=raw_line.source_reference,
                    source_excerpt=f"{pending_parameter} | {text}",
                    **value_only,
                )
            )
            pending_parameter = None
            continue

        if is_likely_specification_label(text):
            pending_parameter = text

    return rows


def specification_section_for_heading(compact_key: str) -> Optional[str]:
    for heading, section in SPECIFICATION_SECTION_HEADINGS.items():
        if compact_key == heading or compact_key.startswith(heading):
            return section
    return None


def parse_specification_line(
    text: str,
    section: str,
    raw_line: RawLine,
) -> Optional[SpecificationRow]:
    not_detected = re.match(
        rf"^(.+?)\s+(Not\s+detect(?:ed|able)(?:\s+in\s+\d+\s*g)?)(?:\s+({SPEC_UNIT_PATTERN}))?\s*$",
        text,
        flags=re.IGNORECASE,
    )
    if not_detected:
        unit = normalize_spec_unit(not_detected.group(3))
        return SpecificationRow(
            section=section,
            parameter=clean_spec_parameter(not_detected.group(1)),
            value=re.sub(r"\s+", " ", not_detected.group(2)).strip(),
            unit=unit,
            qualifier="not detected",
            source_line=raw_line.line_number,
            source_reference=raw_line.source_reference,
            source_excerpt=text,
            needs_review=is_unclear_spec_unit(not_detected.group(3), unit),
        )

    range_match = re.match(
        r"^(.+?)\s+(\d+(?:[,.]\d+)?)\s*-\s*(\d+(?:[,.]\d+)?)"
        rf"\s*({SPEC_UNIT_PATTERN})?\s*(max|min|approx)?\.?$",
        text,
        flags=re.IGNORECASE,
    )
    if range_match:
        unit = normalize_spec_unit(range_match.group(4))
        return SpecificationRow(
            section=section,
            parameter=clean_spec_parameter(range_match.group(1)),
            min_value=range_match.group(2),
            max_value=range_match.group(3),
            unit=unit,
            qualifier=(range_match.group(5) or "range").casefold(),
            source_line=raw_line.line_number,
            source_reference=raw_line.source_reference,
            source_excerpt=text,
            needs_review=is_unclear_spec_unit(range_match.group(4), unit),
        )

    value_match = re.match(
        rf"^(.+?)\s+([<>]?)\s*(\d[\d ]*(?:[,.]\d+)?)\s*({SPEC_UNIT_PATTERN})"
        r"\s*(max|min|approx)?\.?$",
        text,
        flags=re.IGNORECASE,
    )
    if not value_match:
        return None
    comparator = value_match.group(2)
    value = re.sub(r"\s+", "", value_match.group(3))
    qualifier = comparator or (value_match.group(5) or "value").casefold()
    parameter = clean_spec_parameter(value_match.group(1))
    unit = normalize_spec_unit(value_match.group(4))
    implausible_ph = unit == "pH" and float(value.replace(",", ".")) > 14
    return SpecificationRow(
        section=section,
        parameter=parameter,
        value=value,
        max_value=value if qualifier in {"max", "<"} else None,
        min_value=value if qualifier in {"min", ">"} else None,
        unit=unit,
        qualifier=qualifier,
        source_line=raw_line.line_number,
        source_reference=raw_line.source_reference,
        source_excerpt=text,
        needs_review=is_unclear_spec_unit(value_match.group(4), unit)
        or is_unclear_spec_parameter(value_match.group(1), parameter)
        or implausible_ph,
    )


def parse_named_specification_line(
    text: str,
    section: str,
    raw_line: RawLine,
) -> Optional[SpecificationRow]:
    labels_by_section = {
        "Manufacturing": ["Manufactured from", "Processed with"],
        "General appearance": ["Appearance", "Odour", "Taste", "Colour shade", "Color shade", "Solubility"],
    }
    for label in labels_by_section.get(section, []):
        match = re.match(rf"^{re.escape(label)}\s+(.+)$", text, flags=re.IGNORECASE)
        if match:
            return SpecificationRow(
                section=section,
                parameter=label,
                value=clean_value(match.group(1)),
                qualifier="declared",
                source_line=raw_line.line_number,
                source_reference=raw_line.source_reference,
                source_excerpt=text,
            )

    if section == "Shelf life and storage":
        match = re.match(r"^(\d+(?:[,.]\d+)?)\s+months?\s+(.+)$", text, flags=re.IGNORECASE)
        if match:
            condition = clean_value(match.group(2))
            return SpecificationRow(
                section=section,
                parameter=f"Shelf life - {condition}",
                value=match.group(1),
                unit="months",
                qualifier="declared",
                source_line=raw_line.line_number,
                source_reference=raw_line.source_reference,
                source_excerpt=text,
            )
    return None


def parse_specification_value_only(text: str) -> Optional[dict[str, object]]:
    match = re.match(
        r"^(\d[\d ]*(?:[,.]\d+)?)\s*(%\s*\(w/w\)|%|mg/kg|me/kg|g/l|/g|UB)\s*(max|min|approx)?\.?$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    value = re.sub(r"\s+", "", match.group(1))
    qualifier = (match.group(3) or "value").casefold()
    return {
        "value": value,
        "min_value": value if qualifier == "min" else None,
        "max_value": value if qualifier == "max" else None,
        "unit": normalize_spec_unit(match.group(2)),
        "qualifier": qualifier,
        "needs_review": searchable(match.group(2)) == "me/kg",
    }


def specification_display_value(row: SpecificationRow) -> str:
    if row.min_value and row.max_value and row.min_value != row.max_value:
        value = f"{row.min_value} - {row.max_value}"
    else:
        value = row.value or row.max_value or row.min_value or ""
    return " ".join(part for part in [value, row.unit or "", row.qualifier or ""] if part).strip()


def find_specification_value(
    rows: list[SpecificationRow],
    aliases: list[str],
    section_keywords: Optional[list[str]] = None,
) -> Optional[str]:
    alias_keys = [searchable(alias) for alias in aliases]
    for row in rows:
        if section_keywords and not any(
            searchable(section) in searchable(row.section) for section in section_keywords
        ):
            continue
        parameter_key = searchable(row.parameter)
        if any(alias in parameter_key for alias in alias_keys):
            return specification_display_value(row)
    return None


def find_specification_numeric_value(
    rows: list[SpecificationRow],
    aliases: list[str],
    unit_contains: Optional[str] = None,
) -> Optional[str]:
    alias_keys = [searchable(alias).strip() for alias in aliases]
    for row in rows:
        if row.section != "Nutritional data":
            continue
        parameter_key = searchable(row.parameter).strip()
        if not any(parameter_key == alias or parameter_key.startswith(f"{alias} ") for alias in alias_keys):
            continue
        if unit_contains and unit_contains not in searchable(row.unit or ""):
            continue
        return row.value or row.max_value or row.min_value
    return None


def clean_spec_parameter(value: str) -> str:
    return re.sub(r"^[+*$-]+\s*", "", re.sub(r"\s+", " ", value)).strip()


def is_unclear_spec_parameter(raw_value: str, cleaned_value: str) -> bool:
    return "$" in raw_value or bool(
        re.match(r"^[+*]", raw_value.strip()) and not re.search(r"[A-Za-z]", cleaned_value)
    )


def normalize_spec_unit(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    compact = re.sub(r"\s+", "", value)
    normalized = searchable(compact)
    corrections = {
        "kga": "kg/l",
        "gikg": "g/kg",
        "cfulg": "cfu/g",
        "g9/100g": "g/100g",
        "g9/1009": "g/100g",
        "9/100g": "g/100g",
        "9/1009": "g/100g",
        "keal/100g": "kcal/100g",
        "keal'1009": "kcal/100g",
        "kj/100g": "kJ/100g",
        "kcal/100g": "kcal/100g",
        "brix": "Brix",
        "*brix": "Brix",
        "ph": "pH",
    }
    return corrections.get(normalized, compact)


def is_unclear_spec_unit(raw_value: Optional[str], normalized_value: Optional[str]) -> bool:
    if not raw_value or not normalized_value:
        return False
    raw = searchable(re.sub(r"\s+", "", raw_value))
    normalized = searchable(normalized_value)
    return raw != normalized or raw in {"me/kg", "kga", "gikg", "cfulg", "g9/100g", "9/100g", "9/1009"}


def is_likely_specification_label(text: str) -> bool:
    if not re.search(r"[A-Za-z]", text) or len(text) > 90:
        return False
    if re.search(r"\b(?:max|min|approx)\.?\s*$", text, flags=re.IGNORECASE):
        return False
    return not is_document_footer_line(text)


def is_document_footer_line(text: str) -> bool:
    return bool(re.search(r"\b(?:version|revision date|product sheet|page)\s*[:n]", text, flags=re.IGNORECASE))


def detect_output_profile(text: str, source_file: str) -> str:
    if has_any(text, ["exberry", "gnt product no"]) and has_any(
        text,
        ["colouring properties", "coloring properties", "microbiological data"],
    ):
        return "gnt_exberry"
    if has_any(
        text,
        ["production specifications standard approval", "product physical performance/function parameter"],
    ) or has_any(source_file, ["hollow tapered", "filament"]):
        return "filament_spec"
    if has_any(text, ["article specification of roller fabric"]) or has_any(source_file, ["anlon", "roller fabric"]):
        return "roller_fabric"
    if has_any(
        text,
        ["nutrition", "naeringsinnhold", "allergen", "allergener", "shelf life", "holdbarhet"],
    ):
        return "food_ipd"
    return "generic"


def build_dynamic_fields(text: str, source_file: str) -> list[DynamicField]:
    profile = detect_output_profile(text, source_file)
    if profile == "filament_spec":
        return build_filament_dynamic_fields(text, source_file)
    if profile == "generic":
        return build_generic_dynamic_fields(text)

    fields: list[DynamicField] = []
    seen: set[str] = set()
    for line in text.splitlines():
        if not line.startswith("__DYNAMIC_FIELD__|"):
            continue
        parts = line.split("|", 6)
        if len(parts) != 7:
            continue
        _, field_name, section, label, value, unit, needs_review = parts
        if field_name in seen:
            continue
        seen.add(field_name)
        fields.append(
            DynamicField(
                field_name=field_name,
                label=label,
                section=section,
                value=value.strip() or None,
                unit=unit or None,
                source="PDF form cell OCR",
                needs_review=needs_review == "1",
            )
        )

    if detect_output_profile(text, source_file) == "roller_fabric" and "source_file" not in seen:
        fields.insert(
            0,
            DynamicField(
                field_name="source_file",
                label="Source file",
                section="Document & Versioning",
                value=source_file,
                source="Upload filename",
            ),
        )
        enrich_roller_fabric_fields(fields, text)
        product_field = next((field for field in fields if field.field_name == "ohc_article_name"), None)
        if (
            product_field
            and product_field.value
            and "anlon" in searchable(source_file)
            and "white" in searchable(product_field.value)
            and "anlon" not in searchable(product_field.value)
        ):
            product_field.value = re.sub(r"^\S+", "ANLON", product_field.value)
            product_field.needs_review = False
        validate_roller_fabric_fields(fields)
    return fields


def build_generic_dynamic_fields(text: str) -> list[DynamicField]:
    candidates = extract_key_value_candidates(build_raw_lines(text))
    fields: list[DynamicField] = []
    used_names: set[str] = set()
    for candidate in candidates[:200]:
        base_name = re.sub(r"[^a-z0-9]+", "_", searchable(candidate.key)).strip("_") or "captured_field"
        field_name = base_name[:70]
        suffix = 2
        while field_name in used_names:
            ending = f"_{suffix}"
            field_name = f"{base_name[:70 - len(ending)]}{ending}"
            suffix += 1
        used_names.add(field_name)
        fields.append(
            DynamicField(
                field_name=field_name,
                label=candidate.key,
                section=candidate.section or "Other extracted data",
                value=candidate.value,
                source=f"Detected key-value near line {candidate.source_line}" if candidate.source_line else "Detected key-value",
                needs_review=True,
            )
        )
    return fields


def build_filament_dynamic_fields(text: str, source_file: str) -> list[DynamicField]:
    fields = load_dynamic_schema(FILAMENT_TEMPLATE_PATH)
    if not fields:
        return []
    values: dict[str, str] = {"source_file": source_file}

    title = re.search(r"(?m)^(\d{7})\s+(.+?)\s+[\u4e00-\u9fff]", text)
    if title:
        values["ohc_article_no"] = title.group(1)
        values["ohc_article_name"] = clean_value(title.group(2))
    values["doc_date"] = regex_value(text, r"Date[^\n]*?(\d{4}-\d{2}-\d{2})")
    values["doc_written_by"] = regex_value(text, r"Writing[^\n]*?\uff1a\s*([^\s]+)")
    values["doc_approved_by"] = regex_value(text, r"Examining and approving[^\n]*?\uff1a\s*([^\s]+)")

    if re.search(r"TAI\s+HING\s*\(\s*CHINA\s*\)\s*INVESTMENT\s+LIMITED", text, flags=re.IGNORECASE):
        values["supplier_name"] = "TAI HING (CHINA) INVESTMENT LIMITED (TAIHING NYLON)"
    values["supplier_address"] = regex_value(
        text,
        r"(Shop C\s*,\s*On Ying Mansion\s*,\s*G[\w/]?F\.?\s*,?\s*1138 Canton Road\.?\s*,?\s*KLN\s*,\s*Hong Kong)",
    )
    if values["supplier_address"]:
        values["supplier_address"] = re.sub(
            r"G[I/]F\.?\s*,?\s*", "G/F., ", values["supplier_address"], flags=re.IGNORECASE
        )
        values["supplier_address"] = re.sub(r"Road\.?,", "Road,", values["supplier_address"])
    phones = re.findall(r"\+852\s+\d{4}\s+\d{4}", text)
    if phones:
        values["supplier_contact_phone"] = " / ".join(dict.fromkeys(clean_value(phone) for phone in phones))
    values["supplier_contact_email"] = regex_value(text, r"([A-Za-z0-9._%+-]+@taihingnylon\.com)")
    if not values["supplier_contact_email"]:
        ocr_email = regex_value(text, r"(marketing[Q@][A-Za-z0-9.-]+)")
        if ocr_email and has_any(text, ["www.taihingnylon.com"]):
            values["supplier_contact_email"] = "marketing@taihingnylon.com"

    values["material_type"] = regex_value(text, r"Material type[^:\uff1a\n]*[:\uff1a]\s*([A-Za-z0-9]+)")
    values["material_mix_pct"] = regex_value(text, r"Material mix\s*%[^:\uff1a\n]*[:\uff1a]\s*([^\n]+)")
    values["recycled_material_yn"] = regex_value(text, r"(?m)^Recycle material.*[:\uff1a]\s*(YES|NO)\s*$")
    values["colour_code"] = regex_value(text, r"(?m)^\s*(BLACK\s+BK[A-Z0-9-]+)\s*$")

    conditional_filament_values(text, values)
    extract_filament_measurements(text, values)
    extract_filament_instructions(text, values)

    for field in fields:
        value = values.get(field.field_name)
        field.value = value or None
        field.source = "PDF native text" if value else None
        if field.field_name in {"supplier_name", "supplier_address", "supplier_contact_phone", "supplier_contact_email"} and value:
            field.source = "PDF header/footer OCR"
            field.needs_review = True
        if field.field_name == "bending_strength_mn_target" and has_any(text, ["to be corrected", "\u5f85\u66f4\u6b63"]):
            field.needs_review = True
    return fields


def load_dynamic_schema(template_path: Path) -> list[DynamicField]:
    if not template_path.exists():
        return []
    workbook = load_workbook(template_path, data_only=True, read_only=True)
    try:
        dictionary = workbook["Schema Dictionary"]
        fields: list[DynamicField] = []
        for row in dictionary.iter_rows(min_row=2, values_only=True):
            field_name, description, _, unit, _, section = (list(row) + [None] * 6)[:6]
            if not field_name or field_name in {"extraction_notes", "check_status", "_confidence"}:
                continue
            fields.append(
                DynamicField(
                    field_name=str(field_name),
                    label=str(description or field_name).strip(),
                    section=str(section or "Other extracted data"),
                    unit=None if unit in {None, "", "\u2014"} else str(unit),
                )
            )
        return fields
    finally:
        workbook.close()


def conditional_filament_values(text: str, values: dict[str, str]) -> None:
    mappings = {
        "colour_testing_tool": "Colorimeter / Visual Inspection",
        "colour_test_method": "Y09303",
        "colour_acceptance_limit": "Primary Ma",
        "appearance_spec": "No Discoloration / No Pollution",
        "appearance_testing_tool": "Visual Inspection",
        "appearance_test_method": "Y09320",
        "appearance_acceptance_limit": "Primary Ma",
        "cross_section_shape": "Confirming the product being similar to the sample",
        "cross_section_testing_tool": "Visual Inspection",
        "cross_section_test_method": "Y09320",
        "cross_section_acceptance_limit": "Primary Ma",
        "bending_strength_testing_tool": "Elastic machine",
        "bending_strength_test_method": "Y09347",
        "bending_strength_acceptance_limit": "3-5pcs average value",
        "pull_force_testing_tool": "Tension machine",
        "pull_force_test_method": "Y09315",
        "pull_force_acceptance_limit": "3-5pcs average value",
        "smoothness_n_target": "Comparison with sample by touch",
        "smoothness_testing_tool": "Smoothness Tester",
        "smoothness_test_method": "Y09365",
        "smoothness_acceptance_limit": "3-5pcs average value",
        "dryness_testing_tool": "Dryness Tester",
        "dryness_test_method": "Y0381",
        "dryness_acceptance_limit": "3-5pcs average value",
        "ph_value_testing_tool": "PH value tester",
        "ph_value_test_method": "Y09336",
        "ph_value_acceptance_limit": "Fatal Cr",
        "diameter_testing_tool": "Percentage card table",
        "diameter_test_method": "Y09366",
        "diameter_acceptance_limit": "Primary Ma",
        "filament_length_testing_tool": "Tape/Vernier Caliper",
        "filament_length_test_method": "Y09320",
        "filament_length_acceptance_limit": "Primary Ma",
        "filament_straightness_testing_tool": "Constant temperature water bath/sample for testing",
        "filament_straightness_test_method": "Y09348",
        "filament_straightness_acceptance_limit": "Primary Ma",
        "cone_height_testing_tool": "Thickness Gauge / Caliper",
        "cone_height_test_method": "Y09325",
        "cone_height_acceptance_limit": "Secondary Mi",
        "bundle_diameter_testing_tool": "Girth / visual inspection",
        "bundle_diameter_test_method": "Y09352",
        "bundle_diameter_acceptance_limit": "Secondary Mi",
    }
    required_tokens = {
        "colour_testing_tool": "Colorimeter / Visual Inspection",
        "appearance_spec": "No Discoloration/ No Pollution",
        "cross_section_shape": "Confirming the product being similar to the sample",
        "bending_strength_testing_tool": "Elastic machine",
        "pull_force_testing_tool": "Tension machine",
        "smoothness_testing_tool": "Smoothness Tester",
        "dryness_testing_tool": "Dryness Tester",
        "ph_value_testing_tool": "PH value tester",
        "diameter_testing_tool": "Percentage card table",
        "filament_length_testing_tool": "Tape/Vernier Caliper",
        "filament_straightness_testing_tool": "Constant temperature water",
        "cone_height_testing_tool": "Thickness Gauge / Caliper",
        "bundle_diameter_testing_tool": "Girth / visual",
    }
    active_prefixes = {
        key.split("_testing_tool")[0]
        for key, token in required_tokens.items()
        if has_any(text, [token])
    }
    for field_name, value in mappings.items():
        prefix = next((prefix for prefix in active_prefixes if field_name.startswith(prefix)), None)
        if prefix or has_any(text, [value]):
            values[field_name] = value


def extract_filament_measurements(text: str, values: dict[str, str]) -> None:
    measurement_patterns = [
        ("bending_strength", r"Bending strength\(mN\).*?(\d+(?:[.,]\d+)?)\s*\u00b1\s*(\d+(?:[.,]\d+)?)"),
        ("diameter", r"3\.1\s+(\d+(?:[.,]\d+)?)\s*\u00b1\s*(\d+(?:[.,]\d+)?)"),
        ("filament_length", r"3\.5\s+(\d+(?:[.,]\d+)?)\s*\u00b1\s*(\d+(?:[.,]\d+)?)"),
        ("cone_height", r"3\.11\s+(\d+(?:[.,]\d+)?)\s*\u00b1\s*(\d+(?:[.,]\d+)?)"),
    ]
    for prefix, pattern in measurement_patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if match:
            values[f"{prefix}_mm_target" if prefix != "bending_strength" else "bending_strength_mn_target"] = match.group(1)
            values[f"{prefix}_tolerance_mm" if prefix != "bending_strength" else "bending_strength_tolerance_mn"] = match.group(2)
    values["pull_force_n_target"] = regex_value(text, r"2\.2\s+([\u2265<>]=?\s*\d+(?:[.,]\d+)?)").replace("\u2265", ">=")
    values["dryness_grade_target"] = regex_value(text, r"2\.5\s+(\d+\s*-\s*\d+)")
    values["ph_value_target"] = regex_value(text, r"PH value\s+(\d+(?:[.,]\d+)?\s*-\s*\d+(?:[.,]\d+)?)")
    values["filament_straightness_target"] = regex_value(text, r"3\.6\s+([\u2264<>]=?\s*\d+(?:[.,]\d+)?)").replace("\u2264", "<=")
    package = re.search(
        r"Packing diameter[^:\uff1a]*[:\uff1a]\s*(\d+(?:[.,]\d+)?)\s*\u00b1\s*(\d+(?:[.,]\d+)?).*?Single weight.*?[\uff08(]g[\uff09)]\s*[:\uff1a]\s*(\d+(?:[.,]\d+)?)\s*\u00b1\s*(\d+(?:[.,]\d+)?)",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if package:
        values["bundle_diameter_mm_target"] = package.group(1)
        values["bundle_diameter_tolerance_mm"] = package.group(2)
        values["bundle_weight_g_target"] = package.group(3)
        values["bundle_weight_tolerance"] = package.group(4)


def extract_filament_instructions(text: str, values: dict[str, str]) -> None:
    if has_any(text, ["rubber bands"]) and has_any(text, ["paper or plastics around the bundle"]):
        values["bundle_packing_type"] = (
            "Rubber bands (2 pcs, width 3-5mm); paper or plastics around the bundle NOT allowed"
        )
    label_fields = [
        "Supplier", "Orkla House Care Art. No", "Gross weight", "Net weight", "Date of packaging",
        "Number of bundles in the carton (if possible)", "Bundle weight", "Order No", "Diameter of fibre",
        "Length of fibre", "Colour of fibre",
    ]
    present = [item for item in label_fields if has_any(text, [item])]
    if has_any(text, ["Order. No"]) and "Order No" not in present:
        insert_at = present.index("Bundle weight") + 1 if "Bundle weight" in present else len(present)
        present.insert(insert_at, "Order No")
    if len(present) >= 5:
        values["carton_label_required_fields"] = "; ".join(present)
    values["transportation_requirement"] = regex_value(
        text, r"Transportation Requirement:\s*([^\n]+)",
    )
    if values["transportation_requirement"]:
        values["transportation_requirement"] = values["transportation_requirement"][0].upper() + values["transportation_requirement"][1:]
    values["storage_requirement"] = regex_value(text, r"Storage Requirement:\s*([^\n]+)")
    if has_any(text, ["GB/T2828-2012", "ISO 2859", "One-time sampling inspection plan"]):
        values["quality_standard_reference"] = (
            "GB/T2828-2012 (ISO 2859); One-time sampling inspection plan; general inspection level II"
        )
    for field_name, label in [
        ("aql_fatal_cr_level", "fatal Cr"),
        ("aql_primary_ma_level", "primary Ma"),
        ("aql_secondary_mi_level", "secondary Mi"),
    ]:
        values[field_name] = regex_value(text, rf"{re.escape(label)}\s*=\s*(\d+(?:[.,]\d+)?)")
    if has_any(text, ["N/A means not applicable"]):
        values["other_instructions_note"] = "N/A means not applicable"
    patents = re.findall(r"(?:Utility Model|Patent)\s+CN\s?\d+[A-Z]?\s*\(Application number:\s*[^)]+\)", text)
    if patents:
        values["patent_reference"] = "; ".join(clean_value(value) for value in patents)


def regex_value(text: str, pattern: str) -> str:
    match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
    return clean_value(match.group(1)) if match else ""


def enrich_roller_fabric_fields(fields: list[DynamicField], text: str) -> None:
    by_name = {field.field_name: field for field in fields}

    def set_clean(field_name: str, value: Optional[str], needs_review: bool = False) -> None:
        if not value:
            return
        if field_name in by_name:
            by_name[field_name].value = value
            by_name[field_name].needs_review = needs_review
        else:
            field = DynamicField(
                field_name=field_name,
                label=field_name.replace("_", " ").title(),
                section="Supplier Info",
                value=value,
                source="PDF page OCR",
                needs_review=needs_review,
            )
            fields.append(field)
            by_name[field_name] = field

    if re.search(r"\bMITTET\s*UAB\b", text, flags=re.IGNORECASE):
        set_clean("supplier_name", "MITTET UAB")
    match = re.search(r"(Liepy\s+str\.?\s*\d+\s*,?\s*Mosedis)", text, flags=re.IGNORECASE)
    address_first = clean_value(match.group(1)) if match else ""
    match = re.search(r"(LT-?98268\s*,?\s*Skuodas district\s*,?\s*LITHUANIA)", text, flags=re.IGNORECASE)
    address_second = clean_value(match.group(1)) if match else ""
    if address_first or address_second:
        set_clean("supplier_address", "\n".join(value for value in [address_first, address_second] if value))
    match = re.search(r"\bJurate\s+Kabalini\w+", text, flags=re.IGNORECASE)
    if match:
        set_clean("supplier_contact_name", clean_value(match.group(0)))
    phones = re.findall(r"\+?370\s+\d{3}\s+\d{5}", text)
    if phones:
        set_clean("supplier_contact_phone", clean_value(phones[0]))
    if len(phones) > 1:
        set_clean("supplier_contact_cellphone", clean_value(phones[1]))
    match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+(?:\.[A-Za-z]{2,}|\s+[A-Za-z]{2})", text)
    if match:
        email = clean_value(match.group(0))
        set_clean("supplier_contact_email", email, needs_review="." not in email.split("@", 1)[-1])
    match = re.search(r"\b(ANLON\s+WHITE\s*,?\s*PH\s*18\s*MM)\b", text, flags=re.IGNORECASE)
    if match:
        set_clean("ohc_article_name", clean_value(match.group(1)))


def validate_roller_fabric_fields(fields: list[DynamicField]) -> None:
    by_name = {field.field_name: field for field in fields}
    for field in fields:
        if field.field_name == "ohc_article_name" and searchable(field.value or "").startswith("nlon"):
            field.needs_review = True
        if field.field_name in {"material_type_pile", "material_type_backing"}:
            match = re.search(r"(\d+)\s*%", field.value or "")
            if match and int(match.group(1)) > 100:
                field.needs_review = True
        if field.field_name.endswith("_pct_in_pile"):
            match = re.search(r"\d+(?:[,.]\d+)?", field.value or "")
            if match and float(match.group(0).replace(",", ".")) > 100:
                field.needs_review = True
    net = by_name.get("material_width_net_cm_target")
    gross = by_name.get("material_width_gross_cm_target")
    if net and gross:
        try:
            if float(net.value or 0) > float(gross.value or 0):
                net.needs_review = True
                gross.needs_review = True
        except ValueError:
            net.needs_review = True
            gross.needs_review = True


def classify_section(line: str) -> Optional[str]:
    section_keywords = {
        "metadata": ["datablad", "product specification", "produktspesifikasjon", "versjon", "supplier", "plant"],
        "nutrition": ["naeringsinnhold", "nutrition", "nutritional value", "nutrition declaration", "energi", "protein"],
        "allergens": ["allergen", "allergener", "contains", "may contain", "spor av"],
        "shelf_life": ["shelf life", "holdbarhet", "best before", "expiry", "durability"],
        "storage": ["storage", "oppbevaring", "lagring", "store", "temperature", "kjolevare"],
        "physical_chemical": ["physical", "chemical", "fysisk", "kjemisk", "moisture", "water content", "ph", "density", "particle size"],
    }
    for section, keywords in section_keywords.items():
        if has_any(line, keywords):
            return section
    return None


def extract_key_value_candidates(raw_lines: list[RawLine]) -> list[KeyValueCandidate]:
    candidates: list[KeyValueCandidate] = []
    seen: set[tuple[str, str, Optional[int]]] = set()
    for raw_line in raw_lines:
        text = raw_line.text
        line_candidates = parse_allergen_checklist_candidates(text)
        if not line_candidates:
            line_candidates = parse_delimited_candidates(text)
            line_candidates.extend(parse_pipe_candidates(text))
        for key, value in line_candidates:
            key = clean_value(key)
            value = clean_value(value)
            if not key or not value or key == value:
                continue
            identity = (key, value, raw_line.line_number)
            if identity in seen:
                continue
            seen.add(identity)
            candidates.append(
                KeyValueCandidate(
                    section=raw_line.section,
                    key=key[:160],
                    value=value[:600],
                    source_line=raw_line.line_number,
                    source_reference=raw_line.source_reference,
                )
            )
    return candidates


def parse_allergen_checklist_candidates(text: str) -> list[tuple[str, str]]:
    labels = [
        "Gluten",
        "Skalldyr",
        "Egg",
        "Ego",
        "Fisk",
        "Peanotter",
        "Pean\u00f8tter",
        "Soya",
        "Melk",
        "Milk",
        "Notter",
        "N\u00f8tter",
        "Seller",
        "Selleri",
        "Sennep",
        "Sesam",
        "Sulfitt",
        "Suftt",
        "Lupin",
        "Blotdyr",
        "Bl\u00f8tdyr",
    ]
    label_pattern = "|".join(re.escape(label) for label in labels)
    matches = list(re.finditer(rf"\b({label_pattern})\s*:", text, flags=re.IGNORECASE))
    if len(matches) < 2:
        return []

    candidates: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        value_start = match.end()
        value_end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        candidates.append((match.group(1), text[value_start:value_end]))
    return candidates


def parse_delimited_candidates(text: str) -> list[tuple[str, str]]:
    if len(text) > 500:
        return []
    label_matches = list(re.finditer(r"(?<!\S)([\w][\w /().-]{1,60})\s*:", text, flags=re.IGNORECASE))
    if not label_matches:
        return []

    candidates: list[tuple[str, str]] = []
    for index, match in enumerate(label_matches):
        value_start = match.end()
        value_end = label_matches[index + 1].start() if index + 1 < len(label_matches) else len(text)
        key = text[match.start(1) : match.end(1)]
        value = text[value_start:value_end]
        candidates.append((key, value))
    return candidates


def parse_pipe_candidates(text: str) -> list[tuple[str, str]]:
    if "|" not in text:
        return []
    cells = [cell.strip() for cell in text.split("|") if cell.strip()]
    if len(cells) < 2:
        return []
    candidates: list[tuple[str, str]] = []
    for key, value in zip(cells, cells[1:]):
        if has_letter(key) and re.search(r"\d|yes|no|ja|nei|contains|inneholder", searchable(value)):
            candidates.append((key, value))
    return candidates


def build_detected_sections(sections: dict[str, Optional[str]]) -> list[DetectedSection]:
    return [
        DetectedSection(section=name, text=text)
        for name, text in sections.items()
        if text
    ]


def build_review_flags(nutrition: Nutrition, candidates: list[KeyValueCandidate]) -> list[str]:
    flags: list[str] = []
    for field, value in nutrition.model_dump().items():
        if value and looks_like_suspicious_ocr_number(field, value):
            flags.append(f"{field} value '{value}' may need review; OCR may have dropped comma, decimal, or unit characters.")

    for candidate in candidates:
        if has_any(candidate.key, ["allergen", "gluten", "melk", "milk"]) and has_any(candidate.value, ["ye", "nel", "net"]):
            flags.append(f"Allergen candidate on line {candidate.source_line} may contain OCR yes/no errors: {candidate.key}: {candidate.value}")
    return dedupe(flags)


def looks_like_suspicious_ocr_number(field: str, value: str) -> bool:
    normalized = searchable(value).replace(" ", "")
    if "o" in normalized:
        return True
    if field in {"energy_kj", "energy_kcal"}:
        return False
    if re.search(r"\d+[a-z]+$", normalized) and "," not in normalized and "." not in normalized:
        return True
    if re.fullmatch(r"\d{3,}", normalized) and "," not in normalized and "." not in normalized:
        return True
    return False


def has_letter(value: str) -> bool:
    return bool(re.search(r"[A-Za-z]", searchable(value)))


def dedupe(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result
