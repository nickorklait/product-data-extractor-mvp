from io import BytesIO
import re
from copy import copy
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ExtractionResult, IPDAttributeRow

PROJECT_ROOT = Path(__file__).resolve().parents[2]
IPD_TEMPLATE_PATH = PROJECT_ROOT / "samples" / "ipd-template.xlsx"
FABRIC_TEMPLATE_PATH = PROJECT_ROOT / "samples" / "fabric-specs-mittet.xlsx"
FILAMENT_TEMPLATE_PATH = PROJECT_ROOT / "samples" / "filament-specs-taihing-nylon.xlsx"
IPD_HEADERS = ["Spec ID from document", "Node", "Attribute", "Attribute description", "Data", "UoM", "Comments"]
IPD_CONSOLIDATED_HEADERS = [
    "Source file",
    "Product name",
    "Supplier",
    "Plant",
    *IPD_HEADERS,
    "Record type",
    "Needs review",
]
EXCEL_CELL_LIMIT = 32_000


def create_excel(result: ExtractionResult) -> BytesIO:
    if result.output_profile == "roller_fabric":
        return create_schema_excel(result, FABRIC_TEMPLATE_PATH, "Fabric_Specs_Mittet")
    if result.output_profile == "filament_spec":
        return create_schema_excel(result, FILAMENT_TEMPLATE_PATH, "Filament_Specs_Taihing_Nylon")

    output = BytesIO()
    data = result.model_dump()

    first_sheet = (
        {"IPD_Template": build_ipd_template_rows(result)}
        if result.output_profile in {"food_ipd", "gnt_exberry"}
        else {"Extracted_Data": normalize_rows(data.get("dynamic_fields", []))}
    )
    sheets: dict[str, list[dict[str, Any]]] = {
        **first_sheet,
        "Summary": flatten_section(data["metadata"]),
        "Nutrition": flatten_section(data["nutrition"]),
        "Allergens": flatten_section(data["allergens"]),
        "Shelf_Life_Storage": flatten_section(data["shelf_life"]) + flatten_section(data["storage"]),
        "Physical_Chemical": flatten_section(data["physical_chemical"]),
        "Specifications": normalize_rows(data.get("specification_rows", [])),
        "Dynamic_Fields": normalize_rows(data.get("dynamic_fields", [])),
        "Key_Value_Candidates": normalize_rows(data.get("key_value_candidates", [])),
        "Detected_Sections": normalize_rows(data.get("detected_sections", [])),
        "Raw_Lines": normalize_rows(data.get("raw_lines", [])),
        "Review_Flags": [{"Review flag": flag} for flag in data.get("review_flags", [])],
        "Raw_Review": [
            {"Field": "extraction_warnings", "Value": "\n".join(data["extraction_warnings"])},
            {"Field": "raw_text_excerpt", "Value": data["raw_text_excerpt"] or ""},
        ],
    }

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]
            format_worksheet(worksheet)

    output.seek(0)
    return output


def create_batch_excel(results: list[ExtractionResult]) -> BytesIO:
    workbook = Workbook()
    index_sheet = workbook.active
    index_sheet.title = "Batch_Index"
    index_sheet.append(["Source file", "Product", "Output profile", "Data sheet", "Warnings"])

    used_names = {index_sheet.title.casefold()}
    for position, result in enumerate(results, start=1):
        individual = load_workbook(create_excel(result))
        source_sheet = individual[individual.sheetnames[0]]
        preferred_name = result.metadata.document_id or result.metadata.product_name or f"Document {position}"
        sheet_name = unique_sheet_name(str(preferred_name), used_names)
        target_sheet = workbook.create_sheet(sheet_name)
        copy_worksheet(source_sheet, target_sheet)
        index_sheet.append(
            [
                result.metadata.source_file or "",
                result.metadata.product_name or "",
                result.output_profile,
                sheet_name,
                "\n".join(result.extraction_warnings),
            ]
        )
        individual.close()

    format_worksheet(index_sheet)
    index_sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def create_ipd_consolidated_excel(results: list[ExtractionResult]) -> BytesIO:
    """Create a clean overview workbook with complete migration and review detail."""
    workbook = Workbook()
    overview_rows: list[dict[str, Any]] = []
    allergen_rows: list[dict[str, Any]] = []
    mapped_rows: list[dict[str, Any]] = []
    all_ipd_rows: list[dict[str, Any]] = []
    captured_rows: list[dict[str, Any]] = []
    warning_rows: list[dict[str, Any]] = []
    raw_rows: list[dict[str, Any]] = []
    change_rows: list[dict[str, Any]] = []
    specification_rows: list[dict[str, Any]] = []

    for result in results:
        metadata = result.metadata
        common = {
            "Source file": metadata.source_file or "",
            "Product name": metadata.product_name or "",
            "Supplier": metadata.supplier or "",
            "Plant": metadata.plant or "",
        }
        allergen_data = {
            "Gluten": result.allergens.gluten or "",
            "Crustaceans": result.allergens.crustaceans or "",
            "Eggs": result.allergens.eggs or "",
            "Fish": result.allergens.fish or "",
            "Peanuts": result.allergens.peanuts or "",
            "Soybeans": result.allergens.soybeans or "",
            "Milk": result.allergens.milk or "",
            "Nuts": result.allergens.nuts or "",
            "Celery": result.allergens.celery or "",
            "Mustard": result.allergens.mustard or "",
            "Sesame": result.allergens.sesame or "",
            "Sulphites": result.allergens.sulphites or "",
            "Lupin": result.allergens.lupin or "",
            "Molluscs": result.allergens.molluscs or "",
        }
        captured_allergens = sum(bool(value) for value in allergen_data.values())
        document_ipd_rows = build_ipd_template_rows(result)
        found_count = sum(bool(row.get("Data")) for row in document_ipd_rows)
        approved_count = sum(str(row.get("Approved", "")).casefold() in {"yes", "true"} for row in document_ipd_rows)
        declared_allergens = [
            f"{name}: {value}"
            for name, value in allergen_data.items()
            if value and normalize_yes_no_trace(value) != "No"
        ]

        overview_rows.append(
            {
                **common,
                "Document ID": metadata.document_id or "",
                "Version": metadata.version or "",
                "Approved date": metadata.approved_date or "",
                "Shelf life": result.shelf_life.shelf_life_text
                or combine_value_unit(result.shelf_life.shelf_life_value, result.shelf_life.shelf_life_unit),
                "Storage temperature": result.storage.storage_temperature or "",
                "Storage": result.storage.storage_text or "",
                "Allergens reported": "; ".join(declared_allergens) or "None reported",
                "Allergen fields captured": f"{captured_allergens}/14",
                "IPD fields found": f"{found_count}/{len(document_ipd_rows)}",
                "IPD coverage": f"{round((found_count / len(document_ipd_rows)) * 100, 1) if document_ipd_rows else 0}%",
                "Approved fields": f"{approved_count}/{found_count}",
                "Energy kJ": result.nutrition.energy_kj or "",
                "Energy kcal": result.nutrition.energy_kcal or "",
                "Fat g": result.nutrition.fat_g or "",
                "Carbohydrate g": result.nutrition.carbohydrate_g or "",
                "Protein g": result.nutrition.protein_g or "",
                "Salt g": result.nutrition.salt_g or "",
                "Moisture": result.physical_chemical.moisture or "",
                "pH": result.physical_chemical.ph or "",
                "Density": result.physical_chemical.density or "",
                "Particle size": result.physical_chemical.particle_size or "",
                "Warnings": "\n".join(result.extraction_warnings),
            }
        )
        allergen_rows.append(
            {
                **common,
                "Document ID": metadata.document_id or "",
                **allergen_data,
                "Contains": result.allergens.allergens_contains or "",
                "May contain": result.allergens.allergens_may_contain or "",
                "Statement": result.allergens.allergen_statement or "",
                "Captured": f"{captured_allergens}/14",
            }
        )

        for mapped_row in document_ipd_rows:
            output_row = {
                **common,
                **mapped_row,
                "Needs review": "" if mapped_row.get("Data") else "Yes",
            }
            all_ipd_rows.append(output_row)
            if mapped_row.get("Data"):
                mapped_rows.append(output_row)

        for candidate in result.key_value_candidates:
            if candidate.key.casefold().startswith("allergen value -"):
                continue
            captured_rows.append(
                {
                    **common,
                    "Section": candidate.section or "Unclassified",
                    "Field": candidate.key,
                    "Value": candidate.value,
                    "Source line": candidate.source_line or "",
                    "Source reference": candidate.source_reference or "",
                    "Needs review": "Yes",
                }
            )

        raw_text = result.raw_text_full or result.raw_text_excerpt or ""
        raw_chunks = chunk_excel_text(raw_text)
        for chunk_number, chunk in enumerate(raw_chunks, start=1):
            raw_rows.append(
                {
                    **common,
                    "Part": f"{chunk_number}/{len(raw_chunks)}",
                    "Extracted text": chunk,
                }
            )

        for warning in result.extraction_warnings:
            warning_rows.append({**common, "Warning": warning})

        for row in result.specification_rows:
            specification_rows.append(
                {
                    **common,
                    "Section": row.section,
                    "Parameter": row.parameter,
                    "Value": row.value or "",
                    "Min": row.min_value or "",
                    "Max": row.max_value or "",
                    "Unit": row.unit or "",
                    "Qualifier": row.qualifier or "",
                    "Source": row.source_reference or "",
                    "Source excerpt": row.source_excerpt or "",
                    "Needs review": "Yes" if row.needs_review else "",
                }
            )

        for entry in result.change_log:
            change_rows.append({**common, **entry.model_dump()})

    sheets = {
        "Overview": overview_rows,
        "Allergens": allergen_rows,
        "Specifications": specification_rows,
        "IPD_Mapped": mapped_rows,
        "IPD_All_Attributes": all_ipd_rows,
        "Captured_Review": captured_rows,
        "Warnings": warning_rows,
        "Change_Log": change_rows,
        "Raw_Review": raw_rows,
    }
    for index, (sheet_name, rows) in enumerate(sheets.items()):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = sheet_name
        write_rows(worksheet, rows)
        format_worksheet(worksheet)
        worksheet.sheet_view.showGridLines = False

    workbook["Overview"].column_dimensions["A"].width = 46
    workbook["Overview"].column_dimensions["B"].width = 34
    workbook["Allergens"].column_dimensions["A"].width = 46
    workbook["Raw_Review"].column_dimensions["F"].width = 80
    format_allergen_sheet(workbook["Allergens"])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def chunk_excel_text(value: str) -> list[str]:
    if not value:
        return []
    return [value[start:start + EXCEL_CELL_LIMIT] for start in range(0, len(value), EXCEL_CELL_LIMIT)]


def format_allergen_sheet(worksheet) -> None:
    headers = {cell.value: cell.column for cell in worksheet[1]}
    allergen_headers = [
        "Gluten", "Crustaceans", "Eggs", "Fish", "Peanuts", "Soybeans", "Milk",
        "Nuts", "Celery", "Mustard", "Sesame", "Sulphites", "Lupin", "Molluscs",
    ]
    fills = {
        "yes": PatternFill("solid", fgColor="F8D7DA"),
        "no": PatternFill("solid", fgColor="DCEFE5"),
        "other": PatternFill("solid", fgColor="FFF0C7"),
        "missing": PatternFill("solid", fgColor="F1F3F5"),
    }
    for header in allergen_headers:
        column = headers.get(header)
        if not column:
            continue
        worksheet.column_dimensions[get_column_letter(column)].width = max(len(header) + 2, 12)
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row, column)
            normalized = normalize_yes_no_trace(str(cell.value or "")).casefold()
            if normalized == "yes":
                cell.fill = fills["yes"]
            elif normalized == "no":
                cell.fill = fills["no"]
            elif normalized:
                cell.fill = fills["other"]
            else:
                cell.fill = fills["missing"]


def unique_sheet_name(preferred_name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "-", preferred_name).strip() or "Document"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used_names:
        ending = f"-{suffix}"
        candidate = f"{cleaned[:31 - len(ending)]}{ending}"
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def copy_worksheet(source, target) -> None:
    for row in source.iter_rows():
        for cell in row:
            target_cell = target[cell.coordinate]
            target_cell.value = cell.value
            if cell.has_style:
                target_cell.font = copy(cell.font)
                target_cell.fill = copy(cell.fill)
                target_cell.border = copy(cell.border)
                target_cell.alignment = copy(cell.alignment)
                target_cell.number_format = cell.number_format
                target_cell.protection = copy(cell.protection)
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    target.freeze_panes = source.freeze_panes
    target.auto_filter.ref = source.auto_filter.ref
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def create_schema_excel(result: ExtractionResult, template_path: Path, data_sheet_name: str) -> BytesIO:
    if template_path.exists():
        workbook = load_workbook(template_path)
        worksheet = workbook[data_sheet_name]
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Fabric_Specs"
        worksheet.append([field.field_name for field in result.dynamic_fields])

    dynamic_values = {
        field.field_name: field.value or ""
        for field in result.dynamic_fields
    }
    headers = [worksheet.cell(2, column).value for column in range(1, worksheet.max_column + 1)]
    for column, header in enumerate(headers, start=1):
        value = dynamic_values.get(str(header), "")
        if header == "source_file":
            value = result.metadata.source_file or value
        elif header == "extraction_notes":
            value = "\n".join(result.extraction_warnings)
        elif header == "check_status":
            value = "Review" if result.extraction_warnings else "OK"
        elif header == "_confidence":
            value = ""
        worksheet.cell(3, column).value = value

    review_sheets = {
        "Dynamic_Fields": normalize_rows([field.model_dump() for field in result.dynamic_fields]),
        "Specifications": normalize_rows([row.model_dump() for row in result.specification_rows]),
        "Key_Value_Candidates": normalize_rows([item.model_dump() for item in result.key_value_candidates]),
        "Raw_Lines": normalize_rows([item.model_dump() for item in result.raw_lines]),
        "Review_Flags": [{"Review flag": flag} for flag in result.review_flags],
        "Raw_Review": [
            {"Field": "extraction_warnings", "Value": "\n".join(result.extraction_warnings)},
            {"Field": "raw_text_excerpt", "Value": result.raw_text_excerpt or ""},
        ],
    }
    for sheet_name, rows in review_sheets.items():
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
        review_sheet = workbook.create_sheet(sheet_name)
        write_rows(review_sheet, rows)
        format_worksheet(review_sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def write_rows(worksheet, rows: list[dict[str, Any]]) -> None:
    if not rows:
        worksheet.append(["Value"])
        return
    headers = list(rows[0].keys())
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])


def flatten_section(section: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"Field": key, "Value": "" if value is None else value} for key, value in section.items()]


def normalize_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for row in rows:
        normalized.append({key: "" if value is None else value for key, value in row.items()})
    return normalized


def build_ipd_template_rows(result: ExtractionResult) -> list[dict[str, Any]]:
    if result.ipd_rows:
        return [ipd_model_to_row(row) for row in result.ipd_rows]

    template_rows = load_ipd_template_rows()
    metadata = result.metadata
    spec_id = metadata.document_id or metadata.source_file or ""
    candidates = result.key_value_candidates or []

    rows: list[dict[str, Any]] = []
    for template_row in template_rows:
        row = {header: template_row.get(header, "") or "" for header in IPD_HEADERS}
        node = str(row["Node"])
        attribute = str(row["Attribute"])
        description = str(row["Attribute description"])

        data_value, uom, comment = map_ipd_value(result, node, attribute, description, candidates)
        row["Spec ID from document"] = spec_id
        row["Data"] = data_value or ""
        row["UoM"] = uom or row["UoM"] or ""
        row["Comments"] = join_comments(row["Comments"], comment)
        source_reference, source_excerpt = find_ipd_source(
            result,
            node,
            attribute,
            data_value or "",
        )
        row["Source reference"] = source_reference
        row["Source excerpt"] = source_excerpt
        row["Status"] = "Found" if data_value else "Not found"
        row["Approved"] = False
        rows.append(row)

    return rows


def populate_ipd_rows(result: ExtractionResult) -> ExtractionResult:
    if result.ipd_rows:
        return result
    result.ipd_rows = [
        IPDAttributeRow(
            spec_id=row.get("Spec ID from document") or None,
            node=str(row.get("Node") or ""),
            attribute=str(row.get("Attribute") or ""),
            attribute_description=row.get("Attribute description") or None,
            data=row.get("Data") or None,
            uom=row.get("UoM") or None,
            comments=row.get("Comments") or None,
            source_reference=row.get("Source reference") or None,
            source_excerpt=row.get("Source excerpt") or None,
            status=str(row.get("Status") or "Not found"),
            approved=bool(row.get("Approved")),
        )
        for row in build_ipd_template_rows(result)
    ]
    return result


def ipd_model_to_row(row: IPDAttributeRow) -> dict[str, Any]:
    return {
        "Spec ID from document": row.spec_id or "",
        "Node": row.node,
        "Attribute": row.attribute,
        "Attribute description": row.attribute_description or "",
        "Data": row.data or "",
        "UoM": row.uom or "",
        "Comments": row.comments or "",
        "Source reference": row.source_reference or "",
        "Source excerpt": row.source_excerpt or "",
        "Status": row.status,
        "Approved": "Yes" if row.approved else "No",
    }


def find_ipd_source(
    result: ExtractionResult,
    node: str,
    attribute: str,
    value: str,
) -> tuple[str, str]:
    if not value:
        return "Not provided in document", ""

    if simplify(node) == "allergens":
        allergen_marker_names = {
            "gluten": "gluten",
            "crust": "crustaceans",
            "eggs": "eggs",
            "fish": "fish",
            "peanuts": "peanuts",
            "soya": "soybeans",
            "milk": "milk",
            "nuts": "nuts",
            "celery": "celery",
            "mustard": "mustard",
            "sesame": "sesame",
            "so2": "sulphites",
            "lupin": "lupin",
            "molluscs": "molluscs",
        }
        marker_name = allergen_marker_names.get(simplify(attribute))
        marker = f"allergen value - {marker_name}:" if marker_name else ""
        for line in result.raw_lines:
            if marker and marker in line.text.casefold():
                reference = line.source_reference or f"Extracted line {line.line_number}"
                return reference, line.text

    attribute_key = simplify(attribute)
    value_key = simplify(value)
    for line in result.raw_lines:
        line_key = simplify(line.text)
        attribute_matches = attribute_key and (
            attribute_key in line_key or line_key in attribute_key
        )
        value_matches = value_key and value_key in line_key
        if attribute_matches and value_matches:
            reference = line.source_reference or f"Extracted line {line.line_number}"
            return reference, line.text[:600]

    node_section = simplify(node)
    section_aliases = {
        "allergens": "allergens",
        "nutrientsqualityinterfaceper100gm": "nutrition",
        "physicalchemicalproperties": "physical_chemical",
    }
    expected_section = section_aliases.get(node_section)
    for line in result.raw_lines:
        if expected_section and line.section != expected_section:
            continue
        if value_key and value_key in simplify(line.text):
            reference = line.source_reference or f"Extracted line {line.line_number}"
            return reference, line.text[:600]

    return result.metadata.source_file or "Source document", "Mapped from extracted document value."


def load_ipd_template_rows() -> list[dict[str, Any]]:
    if not IPD_TEMPLATE_PATH.exists():
        return fallback_ipd_template_rows()

    workbook = load_workbook(IPD_TEMPLATE_PATH, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    header_row = None
    for row_number in range(1, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, 8)]
        if values[:7] == IPD_HEADERS:
            header_row = row_number
            break

    if header_row is None:
        return fallback_ipd_template_rows()

    rows: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, 8)]
        if not any(value is not None for value in values):
            continue
        rows.append(dict(zip(IPD_HEADERS, values)))
    return rows


def fallback_ipd_template_rows() -> list[dict[str, Any]]:
    basics = [
        ("General", "General description", "General description"),
        ("General", "Raw material name", "Raw material name"),
        ("General", "Country of origin", "Country of origin"),
        ("General", "Shelf life", "Shelf life"),
        ("General", "Storage", "Min"),
        ("General", "Storage", "Max"),
        ("General", "Storage", "Storage conditions"),
        ("Physical & chemical properties", "Density", "Value"),
        ("Physical & chemical properties", "pH", "Min"),
        ("Physical & chemical properties", "pH", "Max"),
        ("Physical & chemical properties", "Water content(%)", "Min"),
        ("Physical & chemical properties", "Water content(%)", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Energy kJ", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Energy kJ", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Energy kcal", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Energy kcal", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Fat", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Fat", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Saturates", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Saturates", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Carbohydrates", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Carbohydrates", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Sugars", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Sugars", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Fibre", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Fibre", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Protien", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Protien", "Max"),
        ("Nutrients Quality interface(per 100 gm)", "Salt", "Min"),
        ("Nutrients Quality interface(per 100 gm)", "Salt", "Max"),
        ("Allergens", "GLUTEN", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "CRUST", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "EGGS", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "FISH", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "PEANUTS", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "SOYA", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "MILK", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "NUTS", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "CELERY", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "MUSTARD", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "SESAME", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "SO2", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "LUPIN", "Yes/No/traces off/traces of quantity"),
        ("Allergens", "MOLLUSCS", "Yes/No/traces off/traces of quantity"),
    ]
    return [
        {
            "Spec ID from document": "",
            "Node": node,
            "Attribute": attribute,
            "Attribute description": description,
            "Data": "",
            "UoM": "",
            "Comments": "",
        }
        for node, attribute, description in basics
    ]


def map_ipd_value(
    result: ExtractionResult,
    node: str,
    attribute: str,
    description: str,
    candidates: list[Any],
) -> tuple[str, str, str]:
    node_key = simplify(node)
    attribute_key = simplify(attribute)
    description_key = simplify(description)

    if node_key == "general":
        mapped = map_general_value(result, attribute_key, description_key, candidates)
        if result.output_profile == "gnt_exberry":
            return mapped
        return mapped if mapped[0] else map_generic_ipd_candidate(attribute, candidates, allow_partial=False)
    if node_key == "physicalchemicalproperties":
        mapped = map_physical_value(result, attribute_key, description_key, candidates)
        if result.output_profile == "gnt_exberry":
            return mapped
        return mapped if mapped[0] else map_generic_ipd_candidate(attribute, candidates, allow_partial=False)
    if node_key == "nutrientsqualityinterfaceper100gm":
        mapped = map_nutrition_value(result, attribute_key, description_key)
        if result.output_profile == "gnt_exberry":
            return mapped
        return mapped if mapped[0] else map_generic_ipd_candidate(attribute, candidates, allow_partial=False)
    if node_key == "allergens":
        return map_allergen_value(result, attribute_key, candidates)
    if node_key in {"microbiologicalproperties", "microbiologicalrequirements", "microbiology"}:
        mapped = map_microbiological_value(result, attribute_key, description_key)
        if result.output_profile == "gnt_exberry":
            return mapped
        return mapped if mapped[0] else map_generic_ipd_candidate(attribute, candidates)
    if result.output_profile == "gnt_exberry":
        return "", "", ""
    return map_generic_ipd_candidate(attribute, candidates)


def map_generic_ipd_candidate(
    attribute: str,
    candidates: list[Any],
    allow_partial: bool = True,
) -> tuple[str, str, str]:
    attribute_key = simplify(attribute)
    if not attribute_key:
        return "", "", ""

    exact_value = ""
    fallback_value = ""
    for candidate in candidates:
        candidate_key = simplify(getattr(candidate, "key", ""))
        candidate_value = str(getattr(candidate, "value", "") or "").strip()
        if not candidate_key or not candidate_value:
            continue
        if candidate_key == attribute_key:
            exact_value = candidate_value
            break
        if allow_partial and len(attribute_key) >= 5 and attribute_key in candidate_key and not fallback_value:
            fallback_value = candidate_value

    value = exact_value or fallback_value
    if not value:
        return "", "", ""
    data_value, uom = split_value_and_unit(value)
    return data_value, uom, "Mapped from a matching source label. Review before approval."


def map_general_value(
    result: ExtractionResult,
    attribute_key: str,
    description_key: str,
    candidates: list[Any],
) -> tuple[str, str, str]:
    if attribute_key == "generaldescription":
        description = result.physical_chemical.physical_chemical_text or result.raw_text_excerpt or ""
        return description, "", "Free text excerpt. Review and shorten before migration." if description else ""
    if attribute_key == "rawmaterialname":
        return result.metadata.product_name or "", "", "Mapped from extracted product name." if result.metadata.product_name else ""
    if attribute_key == "countryoforigin":
        value = find_candidate_value(candidates, ["country of origin", "origin", "opprinnelse", "land"])
        return value, "", "Mapped from extracted key-value candidates. Review before migration." if value else ""
    if attribute_key == "shelflife":
        value = result.shelf_life.shelf_life_text or combine_value_unit(
            result.shelf_life.shelf_life_value,
            result.shelf_life.shelf_life_unit,
        )
        return value, "", "Mapped from shelf-life extraction." if value else ""
    if attribute_key == "storage":
        min_temp, max_temp = split_temperature_range(result.storage.storage_temperature)
        if description_key == "min":
            return min_temp, "C", "Mapped from storage temperature. Review unit and range." if min_temp else ""
        if description_key == "max":
            return max_temp, "C", "Mapped from storage temperature. Review unit and range." if max_temp else ""
        if description_key == "storageconditions":
            return result.storage.storage_text or "", "", "Mapped from storage extraction." if result.storage.storage_text else ""
    return "", "", ""


def map_physical_value(
    result: ExtractionResult,
    attribute_key: str,
    description_key: str,
    candidates: list[Any],
) -> tuple[str, str, str]:
    if attribute_key == "density":
        if description_key != "value":
            return "", "", ""
        row = find_specification_row(
            result,
            ["poured bulk density", "bulk density", "density"],
            ["physical and chemical", "typical values"],
        )
        if row:
            value = specification_row_value(row)
            return value, row.unit or "", "Mapped from the specification table."
        value, uom = split_value_and_unit(result.physical_chemical.density)
        return value, uom, "Mapped from density extraction." if value else ""
    if attribute_key == "ph" and description_key in {"min", "max"}:
        row = find_specification_row(result, ["ph at", "ph of", "ph"], ["physical and chemical"])
        if row:
            value = row.min_value if description_key == "min" else row.max_value
            return value or "", row.unit or "", "Mapped from the specification table range." if value else ""
        value, uom = split_value_and_unit(result.physical_chemical.ph)
        comment = "Single extracted pH value requires review before range mapping." if value else ""
        return value, uom, comment
    if attribute_key == "acidity" and description_key in {"min", "max"}:
        row = find_specification_row(result, ["total acidity as citric acid"], ["physical and chemical"])
        if row:
            value = row.min_value if description_key == "min" else row.max_value or row.value
            return value or "", row.unit or "", "Mapped from the GNT acidity limit." if value else ""
    if attribute_key == "brix" and description_key in {"min", "max"}:
        row = find_specification_row(result, ["dry matter"], ["physical and chemical"])
        if row and simplify(row.unit) == "brix":
            value = row.min_value if description_key == "min" else row.max_value
            return value or "", row.unit or "Brix", "Mapped from the GNT dry-matter Brix range." if value else ""
    if attribute_key in {"watercontent", "watercontentpercent"} and description_key in {"min", "max"}:
        row = find_specification_row(
            result,
            ["loss on drying", "water content", "moisture"],
            ["physical and chemical"],
        )
        if row:
            value = row.min_value if description_key == "min" else row.max_value or row.value
            return value or "", row.unit or "%", "Loss on drying mapped to water-content limit; review before approval." if value else ""
        value, uom = split_value_and_unit(result.physical_chemical.moisture)
        if description_key == "min" and "max" in simplify(value):
            return "", uom or "%", ""
        comment = "Mapped from moisture extraction. Review limit direction." if value else ""
        return value, uom or "%", comment
    if attribute_key == "drymattercontent" and description_key in {"min", "max"}:
        row = find_specification_row(result, ["dry matter content"], ["physical and chemical"])
        if row:
            value = row.min_value if description_key == "min" else row.max_value or row.value
            return value or "", row.unit or "%", "Mapped from the specification table." if value else ""
    return "", "", ""


def find_specification_row(
    result: ExtractionResult,
    aliases: list[str],
    section_keywords: list[str] | None = None,
):
    alias_keys = [simplify(alias) for alias in aliases]
    for row in result.specification_rows:
        if section_keywords and not any(
            simplify(section) in simplify(row.section) for section in section_keywords
        ):
            continue
        parameter_key = simplify(row.parameter)
        if any(alias == parameter_key or alias in parameter_key for alias in alias_keys):
            return row
    return None


def map_microbiological_value(
    result: ExtractionResult,
    attribute_key: str,
    description_key: str,
) -> tuple[str, str, str]:
    aliases = {
        "totalaerobiccount": [
            "total aerobic microbial count",
            "total aerobic count",
            "total plate count aerobic",
        ],
        "totalplatecount": ["total aerobic microbial count", "total plate count"],
        "yeast": ["total yeasts count", "yeast count", "yeast", "yeasis"],
        "mould": ["total moulds count", "mould count", "mould", "moukds"],
        "ecoli": ["escherichia coli", "e. coli", "e coli"],
        "salmonella": ["salmonella", "sabnonella"],
        "salmonellaspp": ["salmonella", "sabnonella"],
    }
    matching_aliases = next(
        (names for key, names in aliases.items() if key == attribute_key or key in attribute_key),
        None,
    )
    if not matching_aliases:
        return "", "", ""
    row = find_specification_row(result, matching_aliases, ["microbiological"])
    if not row:
        return "", "", ""

    if description_key in {"min", "minimum"}:
        value = row.min_value or ""
    elif description_key in {"max", "maximum", "actionlevel"}:
        value = row.max_value or row.value or ""
    else:
        value = specification_row_value(row)
    if not value:
        return "", "", ""
    return value, row.unit or "", "Mapped from the microbiological specification table."


def specification_row_value(row) -> str:
    if row.min_value and row.max_value and row.min_value != row.max_value:
        return f"{row.min_value} - {row.max_value}"
    return row.value or row.max_value or row.min_value or ""


def map_nutrition_value(result: ExtractionResult, attribute_key: str, description_key: str) -> tuple[str, str, str]:
    nutrition_map = {
        "energykj": (result.nutrition.energy_kj, "kJ"),
        "energykcal": (result.nutrition.energy_kcal, "kcal"),
        "fat": (result.nutrition.fat_g, "g"),
        "saturates": (result.nutrition.saturated_fat_g, "g"),
        "carbohydrates": (result.nutrition.carbohydrate_g, "g"),
        "sugars": (result.nutrition.sugars_g, "g"),
        "fibre": (result.nutrition.fibre_g, "g"),
        "protien": (result.nutrition.protein_g, "g"),
        "protein": (result.nutrition.protein_g, "g"),
        "salt": (result.nutrition.salt_g, "g"),
    }
    mapped = nutrition_map.get(attribute_key)
    if not mapped or description_key not in {"min", "max"}:
        return "", "", ""
    value, default_uom = mapped
    data_value, uom = split_value_and_unit(value)
    comment = "Single declared per-100g value placed in Min/Max. Review if IPD expects a range." if data_value else ""
    return data_value, uom or default_uom, comment


def map_allergen_value(
    result: ExtractionResult,
    attribute_key: str,
    candidates: list[Any],
) -> tuple[str, str, str]:
    allergen_aliases = {
        "gluten": ["gluten"],
        "crust": ["crust", "crustacean", "skalldyr"],
        "eggs": ["eggs", "egg", "ego"],
        "fish": ["fish", "fisk"],
        "peanuts": ["peanuts", "peanotter", "peanut"],
        "soya": ["soya", "soy"],
        "milk": ["milk", "melk"],
        "nuts": ["nuts", "notter"],
        "celery": ["celery", "seller", "selleri"],
        "mustard": ["mustard", "sennep"],
        "sesame": ["sesame", "sesam"],
        "so2": ["so2", "sulphite", "sulfitt", "suftt"],
        "lupin": ["lupin"],
        "molluscs": ["molluscs", "mollusk", "blotdyr"],
    }
    allergen_fields = {
        "gluten": result.allergens.gluten,
        "crust": result.allergens.crustaceans,
        "eggs": result.allergens.eggs,
        "fish": result.allergens.fish,
        "peanuts": result.allergens.peanuts,
        "soya": result.allergens.soybeans,
        "milk": result.allergens.milk,
        "nuts": result.allergens.nuts,
        "celery": result.allergens.celery,
        "mustard": result.allergens.mustard,
        "sesame": result.allergens.sesame,
        "so2": result.allergens.sulphites,
        "lupin": result.allergens.lupin,
        "molluscs": result.allergens.molluscs,
    }
    value = allergen_fields.get(attribute_key) or find_candidate_value(
        candidates,
        allergen_aliases.get(attribute_key, [attribute_key]),
    )
    if not value:
        return "", "", ""
    return normalize_yes_no_trace(value), "", "Mapped from allergen checklist/key-value candidates. Review OCR spelling."


def find_candidate_value(candidates: list[Any], labels: list[str]) -> str:
    label_keys = [simplify(label) for label in labels if label]
    for candidate in candidates:
        key = simplify(getattr(candidate, "key", ""))
        if not key:
            continue
        if any(label_key and (label_key in key or key in label_key) for label_key in label_keys):
            return str(getattr(candidate, "value", "") or "").strip()
    return ""


def split_temperature_range(value: str | None) -> tuple[str, str]:
    if not value:
        return "", ""
    numbers = re.findall(r"[+-]?\d+(?:[,.]\d+)?", str(value))
    if len(numbers) >= 2:
        return numbers[0], numbers[1]
    if len(numbers) == 1:
        return numbers[0], numbers[0]
    return "", ""


def combine_value_unit(value: str | None, unit: str | None) -> str:
    if value and unit:
        return f"{value} {unit}"
    return value or unit or ""


def split_value_and_unit(value: str | None) -> tuple[str, str]:
    if not value:
        return "", ""
    text = str(value).strip()
    match = re.match(r"^([<>]?\s*\d+(?:[,.]\d+)?)\s*([A-Za-z%/0-9^]+)?$", text)
    if match:
        return match.group(1).strip(), (match.group(2) or "").strip()
    return text, ""


def normalize_yes_no_trace(value: str) -> str:
    key = simplify(value)
    if any(token in key for token in ["spor", "trace", "maycontain"]):
        return "Traces"
    if any(token in key for token in ["ja", "yes", "ye"]):
        return "Yes"
    if any(token in key for token in ["nei", "no", "nel", "net"]):
        return "No"
    return value.strip()


def simplify(value: Any) -> str:
    text = str(value or "").casefold()
    replacements = {
        "\u00e6": "ae",
        "\u00f8": "o",
        "\u00e5": "a",
        "\u00b0": "",
        "\u2013": "-",
        "\u2014": "-",
        "%": "percent",
        "&": "",
    }
    for source, replacement in replacements.items():
        text = text.replace(source, replacement)
    return re.sub(r"[^a-z0-9]+", "", text)


def join_comments(existing: Any, addition: str) -> str:
    existing_text = str(existing or "").strip()
    addition_text = str(addition or "").strip()
    if existing_text and addition_text:
        return f"{existing_text} | {addition_text}"
    return existing_text or addition_text


def format_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column_cells in worksheet.columns:
        max_length = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 80))
        worksheet.column_dimensions[get_column_letter(column)].width = max(max_length + 2, 14)
